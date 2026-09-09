"""Finalize a technically reviewed reading edition; no scientific execution.

Internal technical QA is separate from human approval and journal submission.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse
import json
import os
import re
import shutil
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import verify_submission_package as v

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
NAMES = {'main': 'NeuralFoil_Measurement_Correction_Manuscript',
         'supplement': 'NeuralFoil_Measurement_Correction_Supplement'}
REVIEWS = [
    'work/render_main/v4/VISUAL_QA_ROOT_1_12.md',
    'work/scientific_review/VISUAL_MAIN_V5_AND_MATH.md',
    'work/scientific_review/MATH_V5_AUDIT.json',
    'work/render_main/v6/VISUAL_QA_ROOT.md',
    'work/literature/VISUAL_SUPP_V5_1_27.md',
    'work/feature_reproduction/VISUAL_SUPP_V5_28_55.md',
    'work/literature/VISUAL_SUPP_V6_CHANGED.md',
    'work/feature_reproduction/VISUAL_SUPP_V6_28_56.md',
    'work/feature_reproduction/V6_BODY_RECONCILIATION.json',
    'work/feature_reproduction/OUTER_PACKAGE_USABILITY_REVIEW.md',
    'work/scientific_review/REFERENCE_REPLAY_REVIEW.md',
    'work/scientific_review/REFERENCE_REPLAY_AUDIT.json',
    'work/scientific_review/ASSEMBLY_V6_AUDIT.json',
    'work/scientific_review/SUBMISSION_PACKAGE_REVIEW.md',
    'work/scientific_review/REVIEW_3.md',
    'work/literature/CLARITY_REVIEW_2.md',
    'work/literature/CITATION_LEDGER.md',
    'work/displays/VISUAL_QA.md',
    'work/render_main/v6/BUILD.json',
    'work/render_main/v6/LAYOUT_AUDIT.json',
    'work/render_main/v6/PAGE_MAP.json',
    'work/render_supplement/v6/BUILD.json',
    'work/render_supplement/v6/LAYOUT_AUDIT.json',
    'work/render_supplement/v6/PAGE_MAP.json',
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', type=Path, default=HERE / 'work/FINAL_TECHNICAL_QA.json')
    args = ap.parse_args()
    output = args.output.resolve()
    if not output.is_relative_to(HERE / 'work'):
        raise ValueError('QA output must be within this successor work directory')
    if output.exists():
        raise FileExistsError('Preserve the previous completed technical record')
    files = {}

    def pin(p, expected=None):
        p = Path(p)
        if p.is_symlink() or not p.is_file() or not p.resolve().is_relative_to(PROJECT):
            raise ValueError('Invalid pin: ' + str(p))
        h = v.sha(p)
        if expected is not None and h != expected:
            raise ValueError('Changed witnessed file: ' + str(p))
        key = Path(os.path.relpath(p, HERE)).as_posix()
        if key in files and files[key] != h:
            raise ValueError('Concurrent input mutation')
        files[key] = h

    documents = {}
    for kind, stem in NAMES.items():
        folder = HERE / 'work' / ('render_' + kind) / 'v6'
        build = json.loads((folder / 'BUILD.json').read_text())
        layout = json.loads((folder / 'LAYOUT_AUDIT.json').read_text())
        if layout['status'] != 'PASS_PROGRAMMATIC' or layout['violations']:
            raise ValueError('Unresolved programmatic layout finding')
        docx = HERE / 'deliverables' / (stem + '.docx')
        pdf = folder / (stem + '.pdf')
        pin(docx, build['docx_sha256'])
        pin(HERE / (kind + '.complete.md'), build['source_sha256'])
        pin(HERE / 'build_documents.py', build['builder_sha256'])
        pin(pdf, layout['pdf_sha256'])
        for row in layout['ledger']:
            pin(folder / ('page-' + str(row['page']) + '.png'), row['png_sha256'])
        target = HERE / 'deliverables' / pdf.name
        if target.exists() and v.sha(target) != v.sha(pdf):
            raise FileExistsError('Refuse to replace a different delivered PDF')
        if not target.exists():
            shutil.copy2(pdf, target)
        pin(target, layout['pdf_sha256'])
        doc = Document(docx)
        if doc.core_properties.author or doc.core_properties.last_modified_by:
            raise ValueError('Unexpected anonymous-document author metadata')
        for tag in ['w:ins', 'w:del', 'w:commentReference', 'w:commentRangeStart']:
            if list(doc.element.iter(qn(tag))):
                raise ValueError('Unresolved revision or comment')
        counts = (len(doc.element.xpath('.//m:oMath')), len(doc.tables), len(doc.inline_shapes))
        if counts != ((84, 3, 3) if kind == 'main' else (134, 28, 1)):
            raise ValueError('Document content counts changed: ' + kind)
        for table in doc.tables:
            if table.rows[0]._tr.find(qn('w:trPr')).find(qn('w:tblHeader')) is None:
                raise ValueError('Missing repeated table header')
        refnums = [int(m.group(1)) for p in doc.paragraphs if (m := re.match(r'^\[(\d+)\]', p.text))]
        if refnums != list(range(1, 18)):
            raise ValueError('Reference sequence changed')
        reader = PdfReader(pdf)
        if reader.metadata.get('/Author', ''):
            raise ValueError('Unexpected PDF author metadata')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        if re.search(r'NAVIGATION_PLACEHOLDER|\bTODO\b|\bTBD\b|\ufffd', text):
            raise ValueError('Unresolved text marker')
        if kind == 'supplement':
            page_map = json.loads((folder / 'PAGE_MAP.json').read_text())
            for title, page in page_map.items():
                if title == 'Reader guide':
                    continue
                if not any(p.text == title + '\t' + str(page) for p in doc.paragraphs):
                    raise ValueError('TOC/page-map mismatch: ' + title)
        documents[kind] = {'pages': len(reader.pages), 'native_math_objects': counts[0],
                           'numbered_display_equations': build['equations'],
                           'tables': counts[1], 'figures': counts[2],
                           'all_pages_visually_accounted_for': True,
                           'author_metadata_empty': True}
    # Reauthenticate the complete, already exposed A/B witness sets. Hash only;
    # no deserialization, model refitting, or unopened outcome collection access.
    witnessed = set()
    for study in ['model_development_20260907_cap_ablation', 'model_development_20260908_adaptive_scale']:
        path = PROJECT / study / 'DELIVERY_QA.json'
        data = json.loads(path.read_text())
        if data['status'] != 'PASS':
            raise ValueError('Study QA did not pass')
        pin(path)
        for name, h in data['source_and_result_sha256'].items():
            pin(Path(name), h)
            witnessed.add(name)
        for p in sorted((PROJECT / study / 'assessment').iterdir()):
            if p.is_file() and p.suffix in ['.json', '.csv']:
                pin(p)
    for relative in ['work/ASSEMBLY.json', 'work/displays/DISPLAY_MANIFEST.json']:
        data = json.loads((HERE / relative).read_text())
        if data['status'] != 'PASS':
            raise ValueError('Assembly/display did not pass')
        pin(HERE / relative)
        for name, h in data['source_sha256'].items():
            pin(PROJECT / name, h)
        for name, h in data.get('outputs', {}).items():
            pin(HERE / name, h)
    for name, expected in {
        '../reproduction_20260908_private/private_reproduction.zip': '7aba2d57fdefe05d942a273630545e964ee996024959997e570017af60295f8a',
        'work/feature_reproduction/feature_reproduction_private.zip': 'eeaa0435dc140d8bf89a76c29fdca188c519aa65dba044d52b190ade3df016b5',
        'work/reference_reproduction/reference_reproduction_private.zip': '5ed12092e3d4f8543791ff750d35f61e6172657d0fc1b1343f6240aaa73fb3bd',
        'work/scientific_review/bounds_portable/bounds_private.zip': '4d3aad28685c65f8b76a5eccb50e10d6209c77a568898844ae270d6c6837ff30',
        'work/scientific_review/EXACT_BOUNDS.json': 'b0e9c6c7b3610505fe1178c78e2fa1077c51619fbe48e1b3dadb174b505a60d7',
        '../paper_revision_20260907/deliverables/NeuralFoil_Measurement_Informed_Drag_Correction.docx': 'e2e289f9f3a11fd94165c8fa50b74f741b1fe650775c88dd71cd298eae49e098',
        '../paper_revision_20260907/deliverables/NeuralFoil_Measurement_Informed_Drag_Correction.pdf': '80abe6bd5061a79daee5e33c23270332cfc9e77dc5bc67610b4b6d71db402c6c',
        '../paper_revision_20260907/deliverables/NeuralFoil_Measurement_Audit_Supplement.docx': '6b2ffdb6603bc4ac1d7ec411fb506f2598e27b5158ae0f85ef0617bddaad01c7',
        '../paper_revision_20260907/deliverables/NeuralFoil_Measurement_Audit_Supplement.pdf': '4625fc18661718c1904e8db1e8b9616deabc558be5cb0ff8cca60551338d4557',
    }.items():
        pin(HERE / name, expected)
    connected = HERE / 'work/reviewer_reproduction'
    pin(connected / 'manifest.json', 'fccad5831d80ce9bc2b3d87e232afc84aac9e1c009f83c7b44ee416886a79528')
    for name, h in json.loads((connected / 'manifest.json').read_text())['files'].items():
        pin(connected / name, h)
    for p in sorted((HERE / 'figures').iterdir()):
        if p.is_file():
            pin(p)
    for relative in ['work/reference_reproduction/release_manifest.json', 'work/scientific_review/bounds_portable/bundle/manifest.json']:
        pin(HERE / relative)
    for relative in REVIEWS:
        pin(HERE / relative)
    for name in ['PACKAGE_README.md', 'verify_submission_package.py', 'test_submission_package.py',
                 'build_submission_package.py', 'final_technical_qa.py',
                 'work/feature_reproduction/NAVIGATION_AND_REPLAY.md']:
        pin(HERE / name)
    for p in (HERE / 'work/literature').iterdir():
        if p.is_file() and p.suffix == '.md':
            pin(p)
    for p in (HERE / 'work/scientific_review/sensitivity_results').iterdir():
        if p.is_file():
            pin(p)
    if json.loads((HERE / 'work/scientific_review/ASSEMBLY_V6_AUDIT.json').read_text())['status'] != 'PASS':
        raise ValueError('Final numerical assembly review did not pass')
    result = {'status': 'PASS_TECHNICAL_PREPARATION',
              'created_utc': datetime.now(timezone.utc).isoformat(),
              'scope': 'Private journal-neutral anonymous reading edition; internal AI-assisted QA, not human approval, public access clearance, journal acceptance, or submission',
              'documents': documents, 'unique_old_study_files_reauthenticated': len(witnessed),
              'files': files, 'review_files': REVIEWS,
              'overnight_minimum_checkpoint_utc': '2026-09-08T12:08:10+00:00',
              'checkpoint_not_claimed_complete': True}
    output.write_text(json.dumps(result, indent=2) + '\n')
    print(json.dumps({k: result[k] for k in ['status', 'documents', 'unique_old_study_files_reauthenticated']}, indent=2))


if __name__ == '__main__':
    main()
