"""Final non-fitting publication QA and package assembly after visual review."""
from pathlib import Path
import csv
import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from assemble_paper import NAMES, num, panel_name

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
ROUND=ROOT/'model_development_20260907_geometry_frontier'

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def rows(p):
    with Path(p).open() as f: return list(csv.DictReader(f))
def clean(s): return re.sub(r'\s+',' ',s).strip()

def main():
    out=HERE/'FINAL_QA.json'
    assert not out.exists(),'Preserve the previous final QA record if a successor is required'
    scientific=json.loads((ROUND/'DELIVERY_QA.json').read_text())
    assert scientific['status']=='PASS'
    for p,h in scientific['authenticated_sha256'].items(): assert sha(p)==h,p
    old=ROOT/'modern_edition'
    oldqa=json.loads((old/'work/visual_qa.json').read_text())
    assert sha(old/'output/pdf/NeuralFoil_Research_Paper.pdf')==oldqa['pdf_sha256']
    assert sha(old/'output/docx/NeuralFoil_Research_Paper.docx')==oldqa['docx_sha256']
    assert sha(old/'manuscript.md')=='7a2d0a6ae9bec2404c49026ace32244bc30beb0e238e9ad5ca7c0b412a12bfbb'
    # The final supplement changes only three pages from the individually reviewed render.
    old_render=HERE/'work/render_supplement_v3'
    final_render=HERE/'work/render_supplement_v5'
    changed=[]
    for i in range(1,74):
        name=f'page-{i}.png'
        if sha(old_render/name)!=sha(final_render/name): changed.append(i)
    assert changed==[4,72,73],changed
    review_paths=[HERE/'work'/f for f in ['visual_review_root.md','visual_review_main_17_32.md',
        'visual_review_supplement_1_36.md','visual_review_supplement_37_73.md']]
    for p in review_paths: assert p.exists() and len(p.read_text())>300,p
    main_doc=HERE/'deliverables/NeuralFoil_Measurement_Informed_Drag_Correction.docx'
    supp_doc=HERE/'deliverables/NeuralFoil_Measurement_Audit_Supplement.docx'
    evidence={}
    pdf_copies=[]
    for kind,docx,render,expected in [('main',main_doc,HERE/'work/render_main_v3',32),
                                    ('supplement',supp_doc,final_render,73)]:
        pdf=render/(docx.stem+'.pdf')
        reader=PdfReader(pdf)
        assert len(reader.pages)==expected
        assert len(list(render.glob('page-*.png')))==expected
        texts=[p.extract_text() for p in reader.pages]
        assert all(len(t)>100 for t in texts)
        assert all('NEURALFOIL MEASUREMENT' in t and 'Kaan Boge' in t for t in texts[1:])
        document=Document(docx)
        assert len(document.sections)==1
        xml=document.element.xml
        assert 'NAVIGATION_PLACEHOLDER' not in xml
        assert not re.search(r'\{\{[A-Z_]+\}\}',xml)
        assert document.core_properties.author=='Kaan Boge'
        anchors={e.get(qn('w:name')) for e in document.element.iter(qn('w:bookmarkStart'))}
        nav=0
        page_map=json.loads((HERE/'work'/kind/'page_map.json').read_text())
        for p in document.paragraphs:
            for link in p._p.iter(qn('w:hyperlink')):
                a=link.get(qn('w:anchor'))
                if a:
                    assert a in anchors,a
                    label=''.join(link.itertext()) if False else ''.join(link.xpath('.//w:t/text()'))
                    if label in page_map:
                        visible=''.join(p._p.xpath('.//w:t/text()'))
                        assert visible.endswith(str(page_map[label])),(label,visible,page_map[label])
                        nav+=1
        assert nav==len(page_map),(kind,nav,page_map)
        dest=HERE/'deliverables'/pdf.name
        assert not dest.exists(),dest
        pdf_copies.append((pdf,dest))
        evidence[kind]={'docx_sha256':sha(docx),'pdf_sha256':sha(pdf),'pages':expected,
            'native_math_nodes':len(list(document.element.iter(qn('m:oMath')))),
            'tables':len(document.tables),'figures':len(document.inline_shapes),'verified_navigation_items':nav,
            'page_png_sha256':{str(i):sha(render/f'page-{i}.png') for i in range(1,expected+1)}}
    md=Document(main_doc)
    sd=Document(supp_doc)
    assert len(md.tables)==10 and len(md.inline_shapes)==2
    assert len(sd.inline_shapes)==11
    panels=rows(ROUND/'assessment/panel_metrics.csv')
    boot=rows(ROUND/'assessment/bootstrap.csv')
    s_tables=[t for t in sd.tables if t.cell(0,0).text.strip()=='View']
    assert len(s_tables)==10
    labels=list(NAMES)
    checked=0
    for label,t in zip(labels,s_tables):
        expected=[p for p in panels if p['candidate']==label]
        assert len(t.rows)==32 and len(expected)==31
        for p,row in zip(expected,t.rows[1:]):
            want=[panel_name(p['panel']),p['rows'],num(p['mae_drag_counts']),num(p['xlarge_CD_improvement_percent']),
                  num(p['mean8_CD_improvement_percent']),num(float(p['p90_absolute_error_CD'])*1e4),num(100*float(p['xlarge_CD_worse_fraction']))]
            assert [clean(c.text) for c in row.cells]==[clean(v) for v in want],(label,p['panel'])
            checked+=1
    t=next(t for t in sd.tables if [c.text for c in t.rows[0].cells]==['Procedure','Assignment','Reference','Reduction %','Conditional 95% interval'])
    assert len(t.rows)==41
    for b,row in zip(boot,t.rows[1:]):
        want=[NAMES[b['candidate']],'A' if int(b['assignment'])==20260906 else 'B',
              'Performance' if b['reference']=='unpenalized_transfer' else 'Half strength',
              num(b['remaining_MAE_reduction_percent']),f"[{num(b['conditional_95pct_lower'])}, {num(b['conditional_95pct_upper'])}]"]
        assert [clean(c.text) for c in row.cells]==[clean(v) for v in want]
    report=json.loads((ROUND/'assessment/report.json').read_text())
    for p in (HERE/'supporting_tables').glob('*.csv'):
        assert sha(p)==report['output_sha256'][str(ROUND/'assessment'/p.name)]
    findings={'status':'PASS','created_utc':datetime.now(timezone.utc).isoformat(),
        'interpretation':'Document and recorded-computation QA; not peer review, public deposition, independent experimental validation or journal acceptance',
        'scientific_delivery_seal_sha256':sha(ROUND/'DELIVERY_QA.json'),
        'scientific_seal_files_reauthenticated':len(scientific['authenticated_sha256']),
        'original_52_page_edition_unchanged':True,'artifacts':evidence,
        'all_105_final_pages_visually_covered':True,'supplement_final_changed_pages_individually_reinspected':changed,
        'other_supplement_pages_pixel_identical_to_reviewed_v3':True,
        'supplement_panel_rows_compared_to_frozen_CSV':checked,'supplement_bootstrap_rows_compared':40,
        'reviews':{p.name:sha(p) for p in review_paths},
        'sources':{str(p.relative_to(HERE)):sha(p) for p in [HERE/'manuscript.md',HERE/'supplement.complete.md',
            HERE/'assemble_paper.py',HERE/'build_reading_edition.py',HERE/'make_figures.py',HERE/'verify_package.py',
            HERE/'REPRODUCIBILITY.md',HERE/'AUTHOR_REVIEW.md',HERE/'CHANGE_RECORD.md']}}
    for source,dest in pdf_copies: shutil.copy2(source,dest)
    out.write_text(json.dumps(findings,indent=2)+'\n')
    # Portable reading package; local scientific directories are intentionally not bundled.
    files=list((HERE/'deliverables').glob('*'))+list((HERE/'supporting_tables').glob('*.csv'))
    files += [HERE/n for n in ['manuscript.md','supplement.complete.md','REPRODUCIBILITY.md','AUTHOR_REVIEW.md','CHANGE_RECORD.md','FINAL_QA.json']]
    files += list((HERE/'figures').glob('*.png'))+list((HERE/'figures').glob('*.svg'))
    files += [HERE/n for n in ['assemble_paper.py','build_reading_edition.py','make_figures.py','inspect_layout.py',
        'verify_package.py','manuscript.template.md','supplement.md','figures/manifest.json',
        'work/EXPERT_COMPONENTS.md','work/METHODS_INTEGRATION_NOTES.md','work/BIBLIOGRAPHY_CHECK.md']]
    files += list((HERE/'blocks').glob('*.md'))
    package=HERE/'NeuralFoil_Research_Paper_20260907.zip'
    assert not package.exists()
    with zipfile.ZipFile(package,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for p in files:z.write(p,str(p.relative_to(HERE)))
    with zipfile.ZipFile(package) as z:
        assert z.testzip() is None
        assert len(z.namelist())==len(files)
    (HERE/'PACKAGE_SHA256.txt').write_text(sha(package)+'  '+package.name+'\n')
    print(json.dumps({'status':'PASS','main_pages':32,'supplement_pages':73,'panel_rows_verified':checked,
        'bootstrap_rows_verified':40,'package_files':len(files),'zip_sha256':sha(package)},indent=2))

if __name__=='__main__':main()
