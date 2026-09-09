"""Bundled-Python DOCX authoring through a tested native OMML Pandoc path.

Creates only this successor edition. Renderer is invoked separately and must
resolve ONLY bundled LibreOffice, never the desktop installation.
"""
from pathlib import Path
import argparse
import importlib.util
import re
import json
import hashlib
import subprocess
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.text.run import Run
from docx.text.paragraph import Paragraph

HERE=Path(__file__).resolve().parent
DESIGN=Path('/PATH_TO_YOUR_HOME/Documents/NeuralFoil_Research_Paper/modern_edition/build_document.py')
spec=importlib.util.spec_from_file_location('preserved_design',DESIGN)
d=importlib.util.module_from_spec(spec);spec.loader.exec_module(d)
TITLE='Accuracy–harm trade-offs in measurement-informed NeuralFoil drag correction'
NAMES={'main':'NeuralFoil_Measurement_Correction_Manuscript','supplement':'NeuralFoil_Measurement_Correction_Supplement'}
TOC_LABELS={
    'S1':'S1 Sources and conservative grouping',
    'S2':'S2 Features and contextual learners',
    'S3':'S3 Strength policies and model comparisons',
    'S4':'S4 Calibration roles and fixed protocols',
    'S5':'S5 Proofs and exact certificates',
    'S6':'S6 Reproduction and numerical environments',
    'S7':'S7 Complete panel results',
    'S8':'S8 Complete interval coverage',
    'S9':'S9 Harm calibration and confidence comparisons',
    'S10':'S10 Measurement label sensitivity',
    'S11':'S11 Claim boundaries and research record',
    'S12':'S12 Refined bounds and predictive utility',
    'S13':'S13 Inference cost and request local reuse',
}

def widths(table,caption):
    head=[c.text.strip() for c in table.rows[0].cells];n=len(head)
    if n==2:return [.7,6.0] if head[0] in ['Position','Index'] else [2.15,4.55]
    if n==3:return [.55,2.8,3.35] if head[0]=='Index' else [2.25,2.225,2.225]
    if head[0]=='Population' and n==4:return [2.4,1.0,1.0,2.3]
    if head==['Procedure','Final t','Historical A / B (%)','W1015-20 all / eligible (%)']:
        return [1.5,.7,2.0,2.5]
    if head==['View','Rows','Generic H','Structural H','Generic KL','Structural KL']:
        return [1.6,.55,1.1375,1.1375,1.1375,1.1375]
    if head==['Context','Identities','Rows','Generic H','Structural H','Generic KL','Structural KL']:
        return [1.6,.65,.6,.95,.95,.95,1.0]
    if head==['Context','B','H t','KL t','Groups','Rows']:
        return [1.7,1.05,1.05,1.05,.85,1.0]
    if head==['View','Rows','H vs XL','H vs mean8','KL vs XL','KL vs mean8']:
        return [1.55,.5,1.1625,1.1625,1.1625,1.1625]
    if head==['Candidate','Reference','Assignment A: point [CI]','Assignment B: point [CI]']:
        return [.95,1.65,2.05,2.05]
    if head==['Route','Warm median [Q25,Q75], ms','Cold request, ms','Cold child process, ms']:
        return [1.4,2.5,1.4,1.4]
    if head==['Route','V6 median, ms','Reuse median, ms','Reuse minus V6, ms','Median reduction, %']:
        return [1.5,1.25,1.25,1.4,1.3]
    if head[0]=='Procedure' and n==5:return [1.9,1.2,1.2,.7,1.8]
    if head[0]=='Population' and n==5:return [1.18,.65,2.1,.87,2.0]
    if head[0]=='Procedure' and n==4:return [2.8,1.5,.95,1.45]
    if head[0]=='View' and n==6:return [2.1,.55,.85,1.05,1.05,1.1]
    if head[0]=='Context' and n==6:return [1.75,.6,.75,1.25,1.1,1.25]
    if head[0]=='Reference' and n==5:return [2.3,.8,1.3,1.15,1.15]
    if head[0]=='View' and n==5:return [2.15,1.14,1.14,1.14,1.13]
    return [6.7/n]*n

def format_doc(doc,kind,pages):
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Inches(.9)
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.75)
    sec.header_distance=sec.footer_distance=Inches(.32)
    sec.different_first_page_header_footer=True
    base=d.style(doc,'Normal',d.BODY,11,after=6);base.paragraph_format.line_spacing=1.22
    for name in ['Body Text','First Paragraph']:
        st=d.style(doc,name,d.BODY,11,after=6);st.base_style=base
        st.paragraph_format.line_spacing=1.22;st.paragraph_format.first_line_indent=Inches(0)
    title_size=25 if kind=='main' else 21
    st=d.style(doc,'Title',d.SANS,title_size,after=16,bold=True)
    st.paragraph_format.line_spacing=1.04;st.paragraph_format.keep_with_next=True
    for name,size,before,after in [('Heading 1',16,18,9),('Heading 2',12.5,14,6),('Heading 3',11.3,11,5)]:
        st=d.style(doc,name,d.SANS,size,before,after,True)
        st.paragraph_format.keep_with_next=True;st.paragraph_format.keep_together=True
    d.style(doc,'Caption',d.SANS,10,4,10)
    d.style(doc,'Compact',d.SANS,10,0,4)
    d.style(doc,'Source Code','Liberation Mono',8.5,0,4)
    for st in doc.styles:
        if st.name=='Verbatim Char':st.font.name='Liberation Mono';st.font.size=Pt(8.5)
    mathpr=doc.settings.element.find(qn('m:mathPr'))
    if mathpr is None:mathpr=d.node('m:mathPr');doc.settings.element.append(mathpr)
    mf=mathpr.find(qn('m:mathFont'))
    if mf is not None:mathpr.remove(mf)
    mathpr.insert(0,d.node('m:mathFont',**{'m:val':d.MATH}))
    # Preserve explicit roman operator/text intent across OMML renderers.
    # Variables without the existing plain-style property are untouched.
    for mr in doc.element.iter(qn('m:r')):
        rpr=mr.find(qn('m:rPr'))
        sty=rpr.find(qn('m:sty')) if rpr is not None else None
        if sty is not None and sty.get(qn('m:val'))=='p':
            if rpr.find(qn('m:nor')) is None:rpr.insert(0,d.node('m:nor'))
            wr=mr.find(qn('w:rPr'))
            if wr is None:wr=d.node('w:rPr');mr.insert(1,wr)
            for name in ('w:i','w:iCs'):
                it=wr.find(qn(name))
                if it is None:it=d.node(name);wr.append(it)
                it.set(qn('w:val'),'0')
    for acc in list(doc.element.iter(qn('m:acc'))):
        char=acc.find(f"{qn('m:accPr')}/{qn('m:chr')}")
        if char is not None and char.get(qn('m:val')) in {'‾','̄'}:
            exp=acc.find(qn('m:e'))
            if exp is not None:
                bar=d.node('m:bar');pr=d.node('m:barPr');pr.append(d.node('m:pos',**{'m:val':'top'}));bar.append(pr);bar.append(deepcopy(exp));acc.getparent().replace(acc,bar)
    # Keep short equation introductions and consecutive display definitions
    # together. Read adjacency before converting display OMML, so a preceding
    # table can never accidentally be treated as the preceding paragraph.
    for p in list(doc.paragraphs):
        if not p._p.xpath('./m:oMathPara'):
            continue
        previous=p._p.getprevious()
        if previous is None or previous.tag!=qn('w:p'):
            continue
        prior=Paragraph(previous,p._parent)
        short_intro=bool(prior.text.strip()) and len(prior.text.strip())<=200
        consecutive_display=bool(previous.xpath('./m:oMathPara'))
        if short_intro or consecutive_display:
            prior.paragraph_format.keep_with_next=True
    headings=[];eqn=0;inrefs=False
    for p in list(doc.paragraphs):
        text=p.text.strip();pf=p.paragraph_format;pf.widow_control=True
        if p.style.name=='Title':
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:d.font(run,d.SANS,title_size,'000000',True)
        if p.style.name.startswith('Heading'):
            for run in p.runs:d.font(run,d.SANS,None,'000000',True)
            anchor=f'nf_section_{len(headings)+1}';d.bookmark(p,anchor,9000+len(headings))
            headings.append({'text':text,'anchor':anchor,'level':int(p.style.name[-1])})
            if text=='References':pf.page_break_before=(kind=='supplement');inrefs=True
            if kind=='supplement' and p.style.name=='Heading 1':pf.page_break_before=True
            if text=='Reader guide':pf.page_break_before=False
            if kind=='supplement' and text.startswith('Procedure '):pf.page_break_before=True
        if text.startswith('Keywords:'):
            for run in p.runs:d.font(run,d.SANS,9.5)
        if text.startswith('Generative AI assistance was used'):
            pf.keep_together=True
        if inrefs and re.match(r'^\[\d+\]',text):
            pf.left_indent=Inches(.25);pf.first_line_indent=Inches(-.25);pf.line_spacing=1.12;pf.space_before=Pt(0);pf.space_after=Pt(7)
            pf.keep_together=True
            # Hyperlink children are not in python-docx's Paragraph.runs.
            # Keep URL and bibliography text at the same readable type size.
            for wr in p._p.xpath('.//w:r'):d.font(Run(wr,p),d.BODY,10)
        if re.match(r'^Table\s+(?:S?\d+)\.',text):
            pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(10);pf.space_after=Pt(5);pf.line_spacing=1.12
            if kind=='supplement' and re.match(r'^Table S(?:28|32|33|34|35|36|37|39|40)\.',text):
                pf.page_break_before=True
            for run in p.runs:d.font(run,d.SANS,10,'000000',True)
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_together=True;pf.keep_with_next=True;pf.line_spacing=1
            pf.space_before=Pt(6);pf.space_after=Pt(3)
        elif re.match(r'^Figure\s+S?\d+\.',text):
            p.style=doc.styles['Caption'];pf.keep_together=True;pf.keep_with_next=False;pf.line_spacing=1.12
            for run in p.runs:d.font(run,d.SANS,10)
        mathparas=p._p.xpath('./m:oMathPara')
        if mathparas:
            eqn+=1;objects=[]
            for mp in mathparas:
                for math in list(mp):
                    if math.tag==qn('m:oMath'):objects.append(deepcopy(math))
                p._p.remove(mp)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(Inches(3.23),WD_TAB_ALIGNMENT.CENTER)
            pf.tab_stops.add_tab_stop(Inches(6.7),WD_TAB_ALIGNMENT.RIGHT)
            p.add_run('\t')
            for math in objects:p._p.append(math)
            d.font(p.add_run(f'\t({"S" if kind=="supplement" else ""}{eqn})'),d.SANS,9.5)
            pf.keep_together=True;pf.space_before=Pt(7);pf.space_after=Pt(9);pf.line_spacing=1.15
        # Prevent a 64-character digest from forcing a protruding unbreakable run.
        for run in p.runs:
            if re.fullmatch(r'[0-9a-f]{64}',run.text):
                run.text=run.text[:32]+'\u200b'+run.text[32:];d.font(run,'Liberation Mono',8.5)
    d.table_widths=widths
    for table in doc.tables:
        d.format_table(table,'')
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):
                # Fit a complete 31-view matrix on one readable page by
                # reducing padding, not type size or scientific content.
                compact_context=(kind=='supplement' and table.cell(0,0).text=='Context' and len(table.rows)==17 and len(table.columns) in (6,7))
                if kind=='supplement' and ((len(table.rows)==32 and len(table.columns) in (5,6)) or compact_context):
                    mar=cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
                    for side in ('top','bottom'):
                        mar.find(qn('w:'+side)).set(qn('w:w'),'40' if compact_context or table.cell(0,1).text=='Eligible rows' else '50')
                for p in cell.paragraphs:
                    keep_whole=(kind=='main' or (kind=='supplement' and (
                        (table.cell(0,0).text=='Reference' and len(table.columns)==5) or
                        (table.cell(0,0).text=='Context' and len(table.rows)==17 and len(table.columns) in (6,7)) or
                        (table.cell(0,0).text=='Route' and len(table.rows)==8 and len(table.columns) in (4,5)) or
                        (len(table.rows)==32 and len(table.columns) in (5,6)))))
                    p.paragraph_format.keep_with_next=(i<len(table.rows)-1) if keep_whole else (i==0)
                    # Keep each candidate's complete comparison block together,
                    # retaining one native table and its repeated column header.
                    # This avoids a lone comparison row on a continuation page.
                    if kind=='supplement' and [c.text.strip() for c in table.rows[0].cells]==['Candidate','Reference','Assignment A: point [CI]','Assignment B: point [CI]']:
                        p.paragraph_format.keep_with_next=(i==0 or (i<len(table.rows)-1 and table.cell(i,0).text==table.cell(i+1,0).text))
                    if j>0 and re.fullmatch(r'[+−\-\d,. /%→]+',p.text.strip()):p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                    for run in p.runs:d.font(run,d.SANS,9.4 if len(table.columns)>=5 else 10,'17212B' if i else '000000',i==0)
    marker=next((p for p in doc.paragraphs if p.text.strip()=='NAVIGATION_PLACEHOLDER'),None)
    if marker:
        for h in headings:
            if h['level']!=1 or not (h['text'].startswith('S') or h['text']=='References'):continue
            p=marker.insert_paragraph_before();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.7),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            label=TOC_LABELS.get(h['text'].split()[0],h['text'])
            d.hyperlink(p,label,h['anchor'],10.5)
            if h['text'] in pages:d.font(p.add_run('\t'+str(pages[h['text']])),d.SANS,10.5)
        marker._p.getparent().remove(marker._p)
    hp=sec.header.paragraphs[0];hp.clear()
    d.font(hp.add_run('NEURALFOIL MEASUREMENT CORRECTION'+('  |  SUPPLEMENT' if kind=='supplement' else '')),d.SANS,8.2,'000000')
    for footer in [sec.footer,sec.first_page_footer]:
        p=footer.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;d.page_field(p)
    sec.first_page_header.paragraphs[0].text=''
    props=doc.core_properties
    props.title=TITLE if kind=='main' else 'Supplementary methods and complete results for measurement informed NeuralFoil drag correction'
    props.author='';props.last_modified_by='';props.comments=''
    props.subject='Anonymous scientific review version with substantive computational assistance disclosure'
    props.keywords='NeuralFoil, airfoil, drag, model discrepancy, grouped evaluation, conformal prediction'
    return headings,eqn

def main():
    ap=argparse.ArgumentParser();ap.add_argument('kind',choices=NAMES);ap.add_argument('--page-map',type=Path);args=ap.parse_args()
    work=HERE/'work'/('render_'+args.kind);work.mkdir(parents=True,exist_ok=True)
    source=HERE/(args.kind+'.complete.md');raw=work/'raw.docx'
    subprocess.run(['/opt/anaconda3/bin/pandoc',str(source),'--from','markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+autolink_bare_uris-implicit_figures',
                    '--to','docx','--resource-path',str(HERE),'--output',str(raw)],check=True,cwd=HERE)
    doc=Document(raw)
    pages=json.loads(args.page_map.read_text()) if args.page_map else {}
    headings,equations=format_doc(doc,args.kind,pages)
    dest=HERE/'deliverables';dest.mkdir(exist_ok=True)
    path=dest/(NAMES[args.kind]+'.docx');doc.save(path)
    report={'status':'BUILT_NOT_YET_VISUALLY_VERIFIED','file':str(path),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
            'builder_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'equations':equations,'all_math_objects':len(doc.element.xpath('.//m:oMath')),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'headings':headings,
            'bundled_libreoffice_only':'/PATH_TO_YOUR_HOME/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice'}
    (work/'BUILD.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({k:report[k] for k in ['status','file','equations','all_math_objects','tables','figures']}))

if __name__=='__main__':main()
