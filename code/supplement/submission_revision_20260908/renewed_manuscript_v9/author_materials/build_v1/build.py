"""Exact reviewed Markdown prose -> restrained editable author materials."""
from datetime import datetime,timezone
import hashlib,io,json,re,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
PINS={'COVER_LETTER_DRAFT':'424c3755ebeaf3da7339c2aca77b49c8b78646ea1dd0e35f3b9ca55f20ad08f2',
      'HIGHLIGHTS_DRAFT':'0fa26a08046d562133be51a64d0ffacdbf3387722b2b1ef5ac9549dc145af4ce',
      'RESEARCH_BRIEF':'424068eba9522b5208be5dc104907e039e7d4305dfcdd91fac708f81393adb79'}
def sha(b):return hashlib.sha256(b).hexdigest()
def blocks(text):
    for b in re.split(r'\n\s*\n',text.strip()):
        if b=='---':continue
        if b.startswith('# '):yield 'Title',b[2:]
        elif b.startswith('## '):yield 'Heading 1',b[3:]
        elif b.startswith('- '):
            for line in b.splitlines():
                assert line.startswith('- ');yield 'List Bullet',line[2:]
        else:yield 'Normal',' '.join(b.splitlines())
def inline(paragraph,text):
    for i,s in enumerate(text.split('**')):
        run=paragraph.add_run(s);run.bold=bool(i%2)
def main():
    out=HERE/'docx';out.mkdir(exist_ok=False)
    manifest=dict(status='DOCX_BUILT_PENDING_RENDER_VISUAL_QA',utc=datetime.now(timezone.utc).isoformat(),
                  builder_sha256=sha(Path(__file__).read_bytes()),sources=PINS,outputs={},text_parity={},python=sys.version)
    for name,h in PINS.items():
        raw=(ROOT/(name+'.md')).read_bytes();assert sha(raw)==h
        entries=list(blocks(raw.decode()))
        doc=Document();sec=doc.sections[0]
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=sec.bottom_margin=Inches(.7)
        sec.left_margin=sec.right_margin=Inches(.8)
        for s in doc.styles:
            if s.type==1:
                s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
        normal=doc.styles['Normal'];normal.font.size=Pt(11)
        normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(6)
        normal.paragraph_format.widow_control=True
        title=doc.styles['Title'];title.font.size=Pt(19);title.font.bold=True
        title.paragraph_format.space_after=Pt(12);title.paragraph_format.keep_with_next=True
        heading=doc.styles['Heading 1'];heading.font.size=Pt(11.5);heading.font.bold=True
        heading.paragraph_format.space_before=Pt(10);heading.paragraph_format.space_after=Pt(5)
        heading.paragraph_format.keep_with_next=True
        for style,text in entries:
            p=doc.add_paragraph(style=style);inline(p,text)
            if style=='List Bullet':p.paragraph_format.space_after=Pt(9)
            if text.startswith(('Unsigned,','Private author-preparation','Internal cold-reader')):
                for run in p.runs:run.font.size=Pt(10)
        props=doc.core_properties;props.author='';props.last_modified_by='';props.title=entries[0][1]
        props.subject='Unsigned author preparation';props.comments=''
        expected=[t.replace('**','') for _,t in entries]
        assert [p.text for p in doc.paragraphs]==expected
        buffer=io.BytesIO();doc.save(buffer);payload=buffer.getvalue()
        with (out/(name+'.docx')).open('xb') as f:f.write(payload)
        reopened=Document(io.BytesIO(payload));assert [p.text for p in reopened.paragraphs]==expected
        assert not reopened.core_properties.author and not reopened.core_properties.last_modified_by
        manifest['outputs'][name+'.docx']=sha(payload)
        manifest['text_parity'][name]=dict(status='EXACT_ORDERED_PROSE',paragraphs=len(expected),
            transformations='Markdown heading/list/emphasis syntax removed; horizontal rules represented by paragraph spacing; all literal prose and punctuation preserved.')
    with (HERE/'BUILD.json').open('x') as f:json.dump(manifest,f,indent=2);f.write('\n')
    print(json.dumps(manifest))
if __name__=='__main__':main()
