"""Explicit-pin technical finalization. No discovery, rendering, or science runs."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile, os
from datetime import datetime, timezone
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

KINDS={'main','supplement'}
ROLES={'assembly','legacy_displays','method_diagram','risk_displays','sensitivity_displays',
       'risk_report','sensitivity_report','risk_portable','sensitivity_portable',
       'environment_receipts','portable_replay_receipts','assembly_audit'}
REQUIRED_MAPS={'assembly':{'/source_sha256'},'legacy_displays':{'/source_sha256','/outputs'},
 'method_diagram':{'/outputs'},'risk_displays':{'/input_sha256','/output_sha256'},
 'sensitivity_displays':{'/input_sha256','/output_sha256'},'risk_report':{'/output_sha256'},
 'sensitivity_report':{'/output_sha256'}}
NEW_ROLES={'qualified_displays','qualified_display_review','qualified_arithmetic_review',
 'qualified_h_report','qualified_kl_report','qualified_h_portable','qualified_kl_portable',
 'qualified_h_replay','qualified_kl_default_replay','qualified_kl_isolated_replay',
 'paired_sole_model_producer','paired_sole_model_replay','paired_sole_model_review'}
ROLES |= NEW_ROLES
ASSERTED_NEW_ROLES=NEW_ROLES-{'qualified_h_portable','qualified_kl_portable'}
REQUIRED_MAPS['qualified_displays']={'/output_sha256'}
REQUIRED_RECORD_LISTS={'qualified_displays':{'/inputs'}}
REQUIRED_SCALAR_PINS={'qualified_displays':{'/source_sha256','/plan_sha256'}}
V9_ROLES={'v8_frozen_qa','v9_assembly','v9_assembly_audit',
 'paired_all_context_producer','paired_all_context_replay','paired_downstream_chain','paired_downstream_review','paired_displays',
 'four_tree_proof_review','four_tree_all_context_producer','four_tree_all_context_replay','four_tree_downstream_chain','four_tree_downstream_review','four_tree_displays',
 'timing_reference_contract','timing_completed_run','timing_saved_record_review','timing_displays',
 'request_local_cache_source_review','request_local_cache_completed_run','request_local_cache_saved_record_review','presentation_block_transform'}
ROLES |= V9_ROLES
REQUIRED_MAPS.update({'v8_frozen_qa':{'/files'},'v9_assembly':{'/inputs','/outputs','/asset_sha256'}})
REQUIRED_SCALAR_PINS['v9_assembly']={'/assembler_sha256','/config/sha256'}
MARKERS=r'@@|NAVIGATION_PLACEHOLDER|<!--|\bTODO\b|\bTBD\b|\ufffd'
def require(ok,message):
    if not ok:raise ValueError(message)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def pointer(d,path):
    for key in path.strip('/').split('/') if path else []:
        key=key.replace('~1','/').replace('~0','~')
        d=d[int(key)] if isinstance(d,list) else d[key]
    return d
class Pins:
    def __init__(self,root):self.root=Path(root).resolve();self.files={}
    def path(self,s):
        p=Path(s);p=p if p.is_absolute() else self.root/p
        require(p.resolve().is_relative_to(self.root),'Path outside project: '+str(p))
        require(not any(x.is_symlink() for x in [p,*p.parents] if x.is_relative_to(self.root)),'Symlink input')
        return p
    def pin(self,spec):
        p=self.path(spec['path']);h=spec['sha256']
        require(re.fullmatch('[0-9a-f]{64}',h) is not None,'Invalid digest')
        require(p.is_file() and sha(p)==h,'Missing or changed pin: '+str(p))
        require(str(p) not in self.files or self.files[str(p)]==h,'Concurrent mutation')
        self.files[str(p)]=h;return p
    def reread(self):
        for p,h in self.files.items():require(sha(Path(p))==h,'Concurrent mutation: '+p)
def visual_page(pins,current,entry):
    """Only actual per-page JSON attestations; no config approval boolean."""
    record_path=pins.pin(entry['ledger']);record=json.loads(record_path.read_text())
    row=pointer(record,entry['record_pointer'])
    require(row['page']==entry['reviewed_page'],'Wrong witnessed page')
    require(row['result']=='PASS' and row['individually_viewed_original_resolution'] is True,'Page not explicitly visually passed')
    report=pins.pin(entry['report']);require(report.stat().st_size>0,'Empty review report')
    reviewed=pins.pin(entry['reviewed_png']);pdf=pins.pin(entry['reviewed_pdf'])
    require(record['pdf_sha256']==sha(pdf),'Review belongs to other PDF')
    require(row['png_sha256']==sha(reviewed),'Review belongs to other PNG')
    require(sha(current)==sha(reviewed),'Changed page needs actual visual review')
    # A differing document/page is accepted only through full PNG byte identity.
    return {'reviewed_page':row['page'],'reviewed_pdf_sha256':sha(pdf),
            'review_report_sha256':sha(report),'ledger_sha256':sha(record_path),
            'lineage':'full_png_byte_identity','png_sha256':sha(current)}
def embedded_fonts(reader):
    seen=set();fonts=[]
    def scan(res):
        if not res:return
        res=res.get_object()
        for name,ref in res.get('/Font',{}).get_object().items() if hasattr(res.get('/Font',{}),'get_object') else res.get('/Font',{}).items():
            f=ref.get_object();ident=(getattr(ref,'idnum',None),str(f.get('/BaseFont')),str(name))
            if ident in seen:continue
            seen.add(ident)
            descendants=f.get('/DescendantFonts',[f])
            for d in descendants:
                d=d.get_object();fd=d.get('/FontDescriptor');fd=fd.get_object() if fd else {}
                ok=any(k in fd for k in ['/FontFile','/FontFile2','/FontFile3'])
                if d.get('/Subtype')=='/Type3':ok=bool(d.get('/CharProcs'))
                require(ok,'Unembedded PDF font: '+str(d.get('/BaseFont',name)))
                fonts.append(str(d.get('/BaseFont',name)))
        for x in res.get('/XObject',{}).get_object().values() if hasattr(res.get('/XObject',{}),'get_object') else res.get('/XObject',{}).values():
            x=x.get_object()
            if x.get('/Subtype')=='/Form':scan(x.get('/Resources'))
    for page in reader.pages:scan(page.get('/Resources'))
    require(bool(fonts),'No fonts found');return fonts
def check_comment_parts(z):
    """Only Word's semantically empty standard comments part is harmless."""
    names=[n for n in z.namelist() if 'comment' in n.lower()]
    require(len(names)<=1 and all(n=='word/comments.xml' for n in names),
            'Unknown or duplicate comment parts')
    for name in names:
        raw=z.read(name)
        require(b'<!DOCTYPE' not in raw.upper() and b'<!ENTITY' not in raw.upper(),
                'Comment DTD or entity declaration')
        try:
            parser=ET.XMLParser(target=ET.TreeBuilder(insert_comments=True,insert_pis=True))
            root=ET.fromstring(raw,parser=parser)
        except ET.ParseError as error:
            raise ValueError('Malformed empty comment part') from error
        require(root.tag==qn('w:comments') and not root.attrib and len(root)==0
                and not (root.text or '').strip() and not (root.tail or '').strip(),
                'Nonempty or nonstandard comment part')

def reading_source_bytes(kind,source,toc_labels,page_map):
    """Exact one-token replacement; deterministic heading order from source."""
    token=b'NAVIGATION_PLACEHOLDER'
    require(isinstance(source,bytes),'Build source bytes required')
    text=source.decode('utf-8')
    if kind!='supplement':
        require(not re.search(MARKERS,text),'Unresolved build-source marker')
        return source
    require(source.count(token)==1 and b'\n'+token+b'\n' in source,
            'Exactly one standalone supplementary navigation token required')
    require(not re.search(MARKERS,text.replace(token.decode(),'')),'Other build-source marker')
    headings=re.findall(r'^# (S[0-9]+ [^\n]+|References)[ \t]*$',text,re.M)
    require(len(headings)==len(set(headings)) and set(headings)==set(toc_labels),'Reading TOC heading mapping mismatch')
    rows=['| Section | PDF page |','| --- | ---: |']
    for heading in headings:
        label=toc_labels[heading];page=page_map[heading]
        require(isinstance(label,str) and label.strip()==label and bool(label) and not any(c in label for c in '|\r\n'),'Invalid reading TOC label')
        require(type(page) is int and page>0,'Invalid reading TOC page')
        rows.append('| '+label+' | '+str(page)+' |')
    result=source.replace(token,'\n'.join(rows).encode('utf-8'))
    require(not re.search(MARKERS,result.decode('utf-8')),'Reading-source marker')
    return result

def check_reading_pair(kind,source,reading,toc_labels,page_map):
    require(reading==reading_source_bytes(kind,source,toc_labels,page_map),'Reading source differs from exact deterministic replacement')

def contract(c):
    require(type(c) is dict and set(c)=={'main','supplement','references'},'Strict content contract fields')
    for k in KINDS:
        require(type(c[k]) is list and len(c[k])==4 and all(type(n) is int and 0<n<10000 for n in c[k]),'Positive bounded native content counts')
    require(type(c['references']) is int and 21<=c['references']<=24,'Explicit reviewed reference contract with at most three appended entries')
    return c

def check_counts(kind,actual,content):
    contract(content)
    require(kind in KINDS and tuple(actual)==tuple(content[kind]) and all(type(n) is int for n in actual),'Content counts mismatch')

def check_references(numbers,total):
    require(type(total) is int and 21<=total<=24,'Reference contract')
    require(numbers==list(range(1,total+1)) and all(type(n) is int for n in numbers),'Reference sequence')

def reference_prefix(old,new,total):
    def entries(text):
        return re.findall(r'^\[(\d+)\] ([^\n]+)$',text,re.M)
    before,after=entries(old),entries(new)
    require([int(n) for n,_ in before]==list(range(1,22)),'Frozen 21-reference source')
    check_references([int(n) for n,_ in after],total)
    require(after[:21]==before,'Original 21 reference entries changed')

def validate_config(c):
    require(c['schema']=='v9-final-qa-1','Config schema')
    require(c.get('finalization_authorized') is True,'Draft config is not authorized for finalization')
    contract(c.get('content_contract'))
    require(set(c['documents'])==KINDS,'Both documents required')
    require(set(c.get('v8_sources',{}))==KINDS,'Explicit frozen V8 reference sources')
    roles=[x['role'] for x in c['evidence']]
    require(ROLES<=set(roles),'Required evidence roles missing')
    require(not any('eight' in x.lower() for x in roles),'Eight-tree evidence not authorized')
    for e in c['evidence']:
        role=e['role']
        require(REQUIRED_MAPS.get(role,set())<={m['pointer'] for m in e.get('maps',[])},'Required manifest map missing: '+role)
        require(REQUIRED_RECORD_LISTS.get(role,set())<={m['pointer'] for m in e.get('record_lists',[])},'Required input-record list missing: '+role)
        require(REQUIRED_SCALAR_PINS.get(role,set())<={m['pointer'] for m in e.get('scalar_pins',[])},'Required scalar source pin missing: '+role)
        if role in ASSERTED_NEW_ROLES:require(bool(e.get('assertions')),'Explicit evidence assertion required: '+role)
        if role in V9_ROLES:
            # Markdown reviews can be directly pinned; structured chain roles
            # must bind selected real maps/scalars and assertions.
            review=role.endswith('_review') or role in {'timing_reference_contract','four_tree_proof_review'}
            if not review:
                require(bool(e.get('assertions')),'Explicit V9 receipt assertion required')
                require(bool(e.get('maps') or e.get('record_lists') or e.get('scalar_pins')),'V9 receipt needs transitive pins')
        if role in {'paired_downstream_chain','four_tree_downstream_chain'}:
            require([x.get('phase') for x in e.get('phase_chain',[])]==['preflight','calibrate','score','assess'],'Explicit complete four-phase chain required')
    for role in V9_ROLES:
        require(roles.count(role)==1,'Unique V9 evidence role')

def frozen_v8(pins,e,d,c):
    require(d.get('status')=='PASS_TECHNICAL_PREPARATION' and d.get('old_scientific_files')==563,'Distinct frozen V8 QA schema')
    require(type(d.get('files')) is dict and bool(d['files']),'Frozen V8 finite inventory')
    spec=next(m for m in e['maps'] if m['pointer']=='/files')
    base=pins.path(spec['base'])
    for k in KINDS:
        source=pins.pin(c['v8_sources'][k])
        rel=Path(os.path.relpath(source,base)).as_posix()
        require(d['files'].get(rel)==sha(source),'V8 reference source absent from frozen QA')
    require(d.get('config_sha256') and d.get('helper_sha256'),'V8 authorizer/helper identities')
    require(d['helper_sha256'] in d['files'].values(),'V8 helper absent from authenticated inventory')
    cfg=pins.path(str(base/d['config_path']))
    require(d['files'].get(Path(os.path.relpath(cfg,base)).as_posix())==d['config_sha256'],'V8 config not bound to inventory')

def phase_chain(pins,rows):
    previous=None
    for row in rows:
        require(set(row)=={'phase','path','sha256','approval','expected_predecessor','output_base'},'Strict selected phase fields')
        record=json.loads(pins.pin(row).read_text());approval=json.loads(pins.pin(row['approval']).read_text())
        pred=row['expected_predecessor']
        require(re.fullmatch('[0-9a-f]{64}',pred) is not None,'Explicit phase predecessor hash')
        require(previous is None or previous==pred,'Phase chain predecessor order')
        require(record.get('status')=='COMPLETE' and record.get('phase')==row['phase'],'Completed matching phase required')
        require(record.get('approval_sha256')==row['approval']['sha256'] and record.get('summary',{}).get('predecessor_sha256')==pred,'Receipt approval/predecessor binding')
        require(approval.get('actual_execution_authorized') is True and approval.get('phase')==row['phase'] and approval.get('registry_sha256')==record.get('registry_sha256') and approval.get('predecessor_sha256')==pred,'Phase authorization binding')
        outputs=record.get('outputs');require(type(outputs) is dict and bool(outputs),'Phase output inventory required')
        for rel,h in outputs.items():pins.pin({'path':str(pins.path(row['output_base'])/rel),'sha256':h})
        previous=row['sha256']

def v9_assembly(pins,e,d,c):
    require(d.get('status')=='ASSEMBLED_NOT_RENDERED_OR_VISUALLY_VERIFIED','Completed V9 assembly required')
    maps={x['pointer']:x for x in e['maps']}
    require(set(d.get('outputs',{}))=={'main.complete.md','supplement.complete.md'},'Exact assembled source outputs')
    base=pins.path(maps['/outputs']['base'])
    for kind in KINDS:
        source=pins.pin(c['documents'][kind]['source'])
        require(base/(kind+'.complete.md')==source,'Assembly output base/document role mismatch')
        require(d['outputs'][kind+'.complete.md']==sha(source),'Assembly output/document source mismatch')
    assets=d.get('asset_sha256',{})
    require(type(assets) is dict and len(assets)==18,'Exact 18 assembly figure assets')
    names=[Path(n) for n in assets];stems={p.stem for p in names}
    require(len(stems)==6 and all({p.suffix for p in names if p.stem==stem}=={'.png','.pdf','.svg'} for stem in stems),'Six figures in three formats')
    require(d.get('counts',{}).get('references_each')==c['content_contract']['references'],'Assembly/reference contract mismatch')


def extra_evidence(pins,e,d):
    # Pointers and bases are explicit configuration; never infer paths from scalar hashes.
    for spec in e.get('scalar_pins',[]):
        pins.pin({'path':spec['path'],'sha256':pointer(d,spec['pointer'])})
    for spec in e.get('record_lists',[]):
        rows=pointer(d,spec['pointer']);require(isinstance(rows,list) and bool(rows),'Empty input-record list')
        for row in rows:
            require(isinstance(row,dict),'Invalid input-record row')
            rel=row['path'];require(isinstance(rel,str),'Input-record path')
            path=Path(rel) if Path(rel).is_absolute() else pins.path(spec['base'])/rel
            pins.pin({'path':str(path),'sha256':row['sha256']})

def document(pins,kind,cfg,content,v8_source):
    require(re.fullmatch(r'v[1-9][0-9]*',cfg['version']) is not None,'Explicit render version required')
    paths={k:pins.pin(cfg[k]) for k in ['source','builder','build','layout','page_map','docx','pdf']}
    for k in ['build','layout','page_map','docx','pdf']:
        require(paths[k].parent.name==cfg['version'],'Render-version path mismatch: '+k)
    build=json.loads(paths['build'].read_text());layout=json.loads(paths['layout'].read_text())
    require(build['source_sha256']==sha(paths['source']) and build['builder_sha256']==sha(paths['builder']) and build['docx_sha256']==sha(paths['docx']),'Build chain mismatch')
    require(layout['status']=='PASS_PROGRAMMATIC' and not layout['violations'] and layout['pdf_sha256']==sha(paths['pdf']),'Layout unresolved')
    doc=Document(paths['docx']);reader=PdfReader(paths['pdf']);page_map=json.loads(paths['page_map'].read_text())
    require(len(reader.pages)==cfg['pages']==layout['pages'],'Page-count mismatch')
    require(not doc.core_properties.author and not doc.core_properties.last_modified_by,'Author metadata')
    require(not (reader.metadata or {}).get('/Author',''),'PDF author metadata')
    for page in reader.pages:
        for ref in page.get('/Annots',[]):
            require(ref.get_object().get('/Subtype') not in ['/Text','/FreeText','/Highlight','/StrikeOut','/Underline','/Squiggly','/FileAttachment'],'PDF review annotation')
    with zipfile.ZipFile(paths['docx']) as z:
        check_comment_parts(z)
    for tag in ['w:ins','w:del','w:moveFrom','w:moveTo','w:commentReference','w:commentRangeStart','w:commentRangeEnd']:
        require(not list(doc.element.iter(qn(tag))),'Revision/comment markup '+tag)
    actual=(build['equations'],len(doc.element.xpath('.//m:oMath')),len(doc.tables),len(doc.inline_shapes))
    check_counts(kind,actual,content)
    numbered=[p for p in doc.paragraphs if p._p.xpath('./m:oMath') and re.search(r'\((?:S)?\d+\)$',p.text)]
    require(len(numbered)==content[kind][0],'Native numbered equations mismatch')
    for t in doc.tables:
        pr=t.rows[0]._tr.find(qn('w:trPr'));header=pr.find(qn('w:tblHeader')) if pr is not None else None
        require(header is not None and header.get(qn('w:val'),'true') not in ['0','false','off'],'Table header not repeated')
    check_references([int(m[1]) for p in doc.paragraphs if (m:=re.match(r'^\[(\d+)\]',p.text))],content['references'])
    reference_prefix(v8_source.read_text(),paths['source'].read_text(),content['references'])
    if kind=='supplement':
        paths['reading_source']=pins.pin(cfg['reading_source'])
        require(paths['reading_source'].parent==paths['source'].parent,'Reading source must preserve local image directory')
        check_reading_pair(kind,paths['source'].read_bytes(),paths['reading_source'].read_bytes(),cfg['toc_labels'],page_map)
    else:
        reading_source_bytes(kind,paths['source'].read_bytes(),{},page_map)
    for txt in [paths.get('reading_source',paths['source']).read_text(),' '.join(doc.element.itertext()),'\n'.join(p.extract_text() or '' for p in reader.pages)]:
        require(not re.search(MARKERS,txt),'Unresolved marker')
    images=[]
    for rel in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',paths['source'].read_text()):
        require(not re.match(r'\w+://',rel),'Remote Markdown image')
        p=pins.path(str(paths['source'].parent/rel.strip('<>')))
        require(str(p) in pins.files,'Markdown image needs explicit evidence pin: '+str(p));images.append(str(p))
    if kind=='supplement':
        headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and (re.match(r'^S\d+ ',p.text) or p.text=='References')]
        require(set(cfg['toc_labels'])==set(headings),'TOC mapping incomplete')
        for heading,label in cfg['toc_labels'].items():
            page=page_map[heading]
            require(any(p.text==label+'\t'+str(page) for p in doc.paragraphs),'TOC number/label mismatch: '+heading)
            actual_text=reader.pages[page-1].extract_text() or ''
            require(re.sub(r'\s+',' ',heading) in re.sub(r'\s+',' ',actual_text),'Heading absent from mapped PDF page')
    require(set(cfg['visual'])=={str(i) for i in range(1,cfg['pages']+1)},'Incomplete visual coverage')
    require({r['page'] for r in layout['ledger']}==set(range(1,cfg['pages']+1)),'Layout ledger page set')
    reviewed=[]
    for row in layout['ledger']:
        p=pins.pin({'path':str(paths['pdf'].parent/f"page-{row['page']}.png"),'sha256':row['png_sha256']})
        reviewed.append({'page':row['page'],**visual_page(pins,p,cfg['visual'][str(row['page'])])})
    return {'version':cfg['version'],'pages':cfg['pages'],'counts':actual,'visual':reviewed,'embedded_fonts':embedded_fonts(reader),'markdown_images':images,
            'artifacts':{k:{'path':str(p),'sha256':sha(p)} for k,p in paths.items()}}
def run(config_path,output,copy_pdfs=False):
    raw=Path(config_path).read_bytes();c=json.loads(raw);validate_config(c)
    pins=Pins(c['project_root']);output=pins.path(output)
    pins.pin({'path':str(Path(config_path).resolve()),'sha256':hashlib.sha256(raw).hexdigest()})
    require(not output.exists(),'Output collision')
    require(set(c['documents'])=={'main','supplement'},'Both documents required')
    require(ROLES<=set(x['role'] for x in c['evidence']),'Required evidence roles missing')
    # All identities and expected statuses are declared; no mutable discovery.
    for e in c['evidence']:
        require(REQUIRED_MAPS.get(e['role'],set())<={m['pointer'] for m in e.get('maps',[])},'Required manifest map missing: '+e['role'])
        if e['role'] in REQUIRED_MAPS or e['role'] in {'assembly_audit','environment_receipts','portable_replay_receipts'}:
            require(bool(e.get('assertions')),'Explicit evidence status assertion required: '+e['role'])
        p=pins.pin(e);d=json.loads(p.read_text()) if e.get('assertions') or e.get('maps') or e.get('record_lists') or e.get('scalar_pins') else None
        extra_evidence(pins,e,d)
        if e['role']=='v8_frozen_qa':frozen_v8(pins,e,d,c)
        if e['role']=='v9_assembly':v9_assembly(pins,e,d,c)
        if e['role'] in {'paired_downstream_chain','four_tree_downstream_chain'}:phase_chain(pins,e['phase_chain'])
        for a in e.get('assertions',[]):require(pointer(d,a['pointer'])==a['equals'],'Evidence assertion '+e['role'])
        for m in e.get('maps',[]):
            values=pointer(d,m['pointer']);require(isinstance(values,dict) and bool(values),'Empty manifest map')
            for rel,h in values.items():
                path=Path(rel) if Path(rel).is_absolute() else pins.path(m['base'])/rel
                pins.pin({'path':str(path),'sha256':h})
    old=pins.pin(c['old_qa']);data=json.loads(old.read_text())
    require(data['status']=='PASS_TECHNICAL_PREPARATION' and data['unique_old_study_files_reauthenticated']==563,'Old QA scope')
    require(len(data['files'])==c['old_qa']['file_count'],'Old QA inventory count')
    for rel,h in data['files'].items():pins.pin({'path':str(pins.path(c['old_qa']['base'])/rel),'sha256':h})
    docs={k:document(pins,k,v,c['content_contract'],pins.pin(c['v8_sources'][k])) for k,v in c['documents'].items()}
    require(output.parent.is_relative_to(pins.path(c['edition'])/'work') or output.parent.is_relative_to(Path(__file__).resolve().parent),'QA output outside edition work or owned helper directory')
    targets=[]
    for d in c['documents'].values():
        if 'deliverable_docx' in d:
            target=pins.path(d['deliverable_docx'])
            require(target.parent==pins.path(c['edition'])/'deliverables','DOCX alias scope')
            pins.pin({'path':str(target),'sha256':d['docx']['sha256']})
    if copy_pdfs:
        for kind,d in c['documents'].items():
            source=pins.pin(d['pdf']);target=pins.path(d['deliverable_pdf'])
            require(target.parent==pins.path(c['edition'])/'deliverables','Delivery scope')
            require(not target.exists() or sha(target)==sha(source),'Refuse differing delivered PDF')
            targets.append((source,target))
    pins.reread();require(Path(config_path).read_bytes()==raw,'Config mutation')
    # Validation is now complete; publication uses exclusive creation only.
    for source,target in targets:
        if not target.exists():
            with target.open('xb') as f:f.write(source.read_bytes())
        pins.pin({'path':str(target),'sha256':sha(source)})
    pins.reread()
    edition=pins.path(c['edition'])
    relative=lambda p:Path(os.path.relpath(p,edition)).as_posix()
    reviews=sorted({relative(pins.path(d['visual'][page][key]['path'])) for d in c['documents'].values() for page in d['visual'] for key in ['report','ledger']})
    result={'status':'PASS_TECHNICAL_PREPARATION','scope':'Internal computational QA only; no author approval, rights clearance, submission, or physical accuracy certification',
      'created_utc':datetime.now(timezone.utc).isoformat(),'config_sha256':hashlib.sha256(raw).hexdigest(),
      'helper_sha256':sha(Path(__file__)),'documents':docs,'old_scientific_files':563,'old_total_files':len(data['files']),
      'files':{relative(p):h for p,h in pins.files.items()},'review_files':reviews,
      'config_path':relative(Path(config_path).resolve()),'copied_or_identical_pdfs':[relative(t) for _,t in targets]}
    with output.open('x') as f:json.dump(result,f,indent=2);f.write('\n')
    return result
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('config',type=Path);ap.add_argument('--output',required=True,type=Path);ap.add_argument('--copy-pdfs',action='store_true');a=ap.parse_args()
    print(json.dumps(run(a.config,a.output,a.copy_pdfs),indent=2))
