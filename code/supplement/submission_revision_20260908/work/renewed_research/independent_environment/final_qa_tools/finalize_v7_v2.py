"""Explicit-pin technical finalization. No discovery, rendering, or science runs."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile, os
from datetime import datetime, timezone
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

COUNTS={'main':(16,116,3,4),'supplement':(28,203,33,1)}
ROLES={'assembly','legacy_displays','method_diagram','risk_displays','sensitivity_displays',
       'risk_report','sensitivity_report','risk_portable','sensitivity_portable',
       'environment_receipts','portable_replay_receipts','assembly_audit'}
REQUIRED_MAPS={'assembly':{'/source_sha256'},'legacy_displays':{'/source_sha256','/outputs'},
 'method_diagram':{'/outputs'},'risk_displays':{'/input_sha256','/output_sha256'},
 'sensitivity_displays':{'/input_sha256','/output_sha256'},'risk_report':{'/output_sha256'},
 'sensitivity_report':{'/output_sha256'}}
MARKERS=r'NAVIGATION_PLACEHOLDER|<!--|\bTODO\b|\bTBD\b|\ufffd'
def require(ok,message):
    if not ok:raise ValueError(message)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def pointer(d,path):
    for key in path.strip('/').split('/') if path else []:
        key=key.replace('~1','/').replace('~0','~')
        d=d[int(key)] if isinstance(d,list) else d[key]
    return d
class Pins:
    def __init__(self,root):self.root=Path(root).resolve();self.files={}
    def path(self,s):
        p=Path(s);p=p if p.is_absolute() else self.root/p
        require(p.resolve().is_relative_to(self.root),'Path outside project: '+str(p))
        require(not any(x.is_symlink() for x in [p,*p.parents] if x.is_relative_to(self.root)),'Symlink input')
        return p
    def pin(self,spec):
        p=self.path(spec['path']);h=spec['sha256']
        require(re.fullmatch('[0-9a-f]{64}',h) is not None,'Invalid digest')
        require(p.is_file() and sha(p)==h,'Missing or changed pin: '+str(p))
        require(str(p) not in self.files or self.files[str(p)]==h,'Concurrent mutation')
        self.files[str(p)]=h;return p
    def reread(self):
        for p,h in self.files.items():require(sha(Path(p))==h,'Concurrent mutation: '+p)
def visual_page(pins,current,entry):
    """Only actual per-page JSON attestations; no config approval boolean."""
    record_path=pins.pin(entry['ledger']);record=json.loads(record_path.read_text())
    row=pointer(record,entry['record_pointer'])
    require(row['page']==entry['reviewed_page'],'Wrong witnessed page')
    require(row['result']=='PASS' and row['individually_viewed_original_resolution'] is True,'Page not explicitly visually passed')
    report=pins.pin(entry['report']);require(report.stat().st_size>0,'Empty review report')
    reviewed=pins.pin(entry['reviewed_png']);pdf=pins.pin(entry['reviewed_pdf'])
    require(record['pdf_sha256']==sha(pdf),'Review belongs to other PDF')
    require(row['png_sha256']==sha(reviewed),'Review belongs to other PNG')
    require(sha(current)==sha(reviewed),'Changed page needs actual visual review')
    # A differing document/page is accepted only through full PNG byte identity.
    return {'reviewed_page':row['page'],'reviewed_pdf_sha256':sha(pdf),
            'review_report_sha256':sha(report),'ledger_sha256':sha(record_path),
            'lineage':'full_png_byte_identity','png_sha256':sha(current)}
def embedded_fonts(reader):
    seen=set();fonts=[]
    def scan(res):
        if not res:return
        res=res.get_object()
        for name,ref in res.get('/Font',{}).get_object().items() if hasattr(res.get('/Font',{}),'get_object') else res.get('/Font',{}).items():
            f=ref.get_object();ident=(getattr(ref,'idnum',None),str(f.get('/BaseFont')),str(name))
            if ident in seen:continue
            seen.add(ident)
            descendants=f.get('/DescendantFonts',[f])
            for d in descendants:
                d=d.get_object();fd=d.get('/FontDescriptor');fd=fd.get_object() if fd else {}
                ok=any(k in fd for k in ['/FontFile','/FontFile2','/FontFile3'])
                if d.get('/Subtype')=='/Type3':ok=bool(d.get('/CharProcs'))
                require(ok,'Unembedded PDF font: '+str(d.get('/BaseFont',name)))
                fonts.append(str(d.get('/BaseFont',name)))
        for x in res.get('/XObject',{}).get_object().values() if hasattr(res.get('/XObject',{}),'get_object') else res.get('/XObject',{}).values():
            x=x.get_object()
            if x.get('/Subtype')=='/Form':scan(x.get('/Resources'))
    for page in reader.pages:scan(page.get('/Resources'))
    require(bool(fonts),'No fonts found');return fonts
def check_comment_parts(z):
    """Only Word's semantically empty standard comments part is harmless."""
    names=[n for n in z.namelist() if 'comment' in n.lower()]
    require(len(names)<=1 and all(n=='word/comments.xml' for n in names),
            'Unknown or duplicate comment parts')
    for name in names:
        raw=z.read(name)
        require(b'<!DOCTYPE' not in raw.upper() and b'<!ENTITY' not in raw.upper(),
                'Comment DTD or entity declaration')
        try:
            parser=ET.XMLParser(target=ET.TreeBuilder(insert_comments=True,insert_pis=True))
            root=ET.fromstring(raw,parser=parser)
        except ET.ParseError as error:
            raise ValueError('Malformed empty comment part') from error
        require(root.tag==qn('w:comments') and not root.attrib and len(root)==0
                and not (root.text or '').strip() and not (root.tail or '').strip(),
                'Nonempty or nonstandard comment part')

def document(pins,kind,cfg):
    require(re.fullmatch(r'v[1-9][0-9]*',cfg['version']) is not None,'Explicit render version required')
    paths={k:pins.pin(cfg[k]) for k in ['source','builder','build','layout','page_map','docx','pdf']}
    for k in ['build','layout','page_map','docx','pdf']:
        require(paths[k].parent.name==cfg['version'],'Render-version path mismatch: '+k)
    build=json.loads(paths['build'].read_text());layout=json.loads(paths['layout'].read_text())
    require(build['source_sha256']==sha(paths['source']) and build['builder_sha256']==sha(paths['builder']) and build['docx_sha256']==sha(paths['docx']),'Build chain mismatch')
    require(layout['status']=='PASS_PROGRAMMATIC' and not layout['violations'] and layout['pdf_sha256']==sha(paths['pdf']),'Layout unresolved')
    doc=Document(paths['docx']);reader=PdfReader(paths['pdf']);page_map=json.loads(paths['page_map'].read_text())
    require(len(reader.pages)==cfg['pages']==layout['pages'],'Page-count mismatch')
    require(not doc.core_properties.author and not doc.core_properties.last_modified_by,'Author metadata')
    require(not (reader.metadata or {}).get('/Author',''),'PDF author metadata')
    for page in reader.pages:
        for ref in page.get('/Annots',[]):
            require(ref.get_object().get('/Subtype') not in ['/Text','/FreeText','/Highlight','/StrikeOut','/Underline','/Squiggly','/FileAttachment'],'PDF review annotation')
    with zipfile.ZipFile(paths['docx']) as z:
        check_comment_parts(z)
    for tag in ['w:ins','w:del','w:moveFrom','w:moveTo','w:commentReference','w:commentRangeStart','w:commentRangeEnd']:
        require(not list(doc.element.iter(qn(tag))),'Revision/comment markup '+tag)
    actual=(build['equations'],len(doc.element.xpath('.//m:oMath')),len(doc.tables),len(doc.inline_shapes))
    require(actual==COUNTS[kind],'Content counts mismatch')
    numbered=[p for p in doc.paragraphs if p._p.xpath('./m:oMath') and re.search(r'\((?:S)?\d+\)$',p.text)]
    require(len(numbered)==COUNTS[kind][0],'Native numbered equations mismatch')
    for t in doc.tables:
        pr=t.rows[0]._tr.find(qn('w:trPr'));header=pr.find(qn('w:tblHeader')) if pr is not None else None
        require(header is not None and header.get(qn('w:val'),'true') not in ['0','false','off'],'Table header not repeated')
    require([int(m[1]) for p in doc.paragraphs if (m:=re.match(r'^\[(\d+)\]',p.text))]==list(range(1,20)),'Reference sequence')
    for txt in [paths['source'].read_text(),' '.join(doc.element.itertext()),'\n'.join(p.extract_text() or '' for p in reader.pages)]:
        require(not re.search(MARKERS,txt),'Unresolved marker')
    images=[]
    for rel in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',paths['source'].read_text()):
        require(not re.match(r'\w+://',rel),'Remote Markdown image')
        p=pins.path(str(paths['source'].parent/rel.strip('<>')))
        require(str(p) in pins.files,'Markdown image needs explicit evidence pin: '+str(p));images.append(str(p))
    if kind=='supplement':
        headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and (re.match(r'^S\d+ ',p.text) or p.text=='References')]
        require(set(cfg['toc_labels'])==set(headings),'TOC mapping incomplete')
        for heading,label in cfg['toc_labels'].items():
            page=page_map[heading]
            require(any(p.text==label+'\t'+str(page) for p in doc.paragraphs),'TOC number/label mismatch: '+heading)
            actual_text=reader.pages[page-1].extract_text() or ''
            require(re.sub(r'\s+',' ',heading) in re.sub(r'\s+',' ',actual_text),'Heading absent from mapped PDF page')
    require(set(cfg['visual'])=={str(i) for i in range(1,cfg['pages']+1)},'Incomplete visual coverage')
    require({r['page'] for r in layout['ledger']}==set(range(1,cfg['pages']+1)),'Layout ledger page set')
    reviewed=[]
    for row in layout['ledger']:
        p=pins.pin({'path':str(paths['pdf'].parent/f"page-{row['page']}.png"),'sha256':row['png_sha256']})
        reviewed.append({'page':row['page'],**visual_page(pins,p,cfg['visual'][str(row['page'])])})
    return {'version':cfg['version'],'pages':cfg['pages'],'counts':actual,'visual':reviewed,'embedded_fonts':embedded_fonts(reader),'markdown_images':images,
            'artifacts':{k:{'path':str(p),'sha256':sha(p)} for k,p in paths.items()}}
def run(config_path,output,copy_pdfs=False):
    raw=Path(config_path).read_bytes();c=json.loads(raw);require(c['schema']=='v7-final-qa-1','Config schema')
    require(c.get('finalization_authorized') is True,'Draft config is not authorized for finalization')
    pins=Pins(c['project_root']);output=pins.path(output)
    pins.pin({'path':str(Path(config_path).resolve()),'sha256':hashlib.sha256(raw).hexdigest()})
    require(not output.exists(),'Output collision')
    require(set(c['documents'])=={'main','supplement'},'Both documents required')
    require(ROLES<=set(x['role'] for x in c['evidence']),'Required evidence roles missing')
    # All identities and expected statuses are declared; no mutable discovery.
    for e in c['evidence']:
        require(REQUIRED_MAPS.get(e['role'],set())<={m['pointer'] for m in e.get('maps',[])},'Required manifest map missing: '+e['role'])
        if e['role'] in REQUIRED_MAPS or e['role'] in {'assembly_audit','environment_receipts','portable_replay_receipts'}:
            require(bool(e.get('assertions')),'Explicit evidence status assertion required: '+e['role'])
        p=pins.pin(e);d=json.loads(p.read_text()) if e.get('assertions') or e.get('maps') else None
        for a in e.get('assertions',[]):require(pointer(d,a['pointer'])==a['equals'],'Evidence assertion '+e['role'])
        for m in e.get('maps',[]):
            values=pointer(d,m['pointer']);require(isinstance(values,dict) and bool(values),'Empty manifest map')
            for rel,h in values.items():
                path=Path(rel) if Path(rel).is_absolute() else pins.path(m['base'])/rel
                pins.pin({'path':str(path),'sha256':h})
    old=pins.pin(c['old_qa']);data=json.loads(old.read_text())
    require(data['status']=='PASS_TECHNICAL_PREPARATION' and data['unique_old_study_files_reauthenticated']==563,'Old QA scope')
    require(len(data['files'])==c['old_qa']['file_count'],'Old QA inventory count')
    for rel,h in data['files'].items():pins.pin({'path':str(pins.path(c['old_qa']['base'])/rel),'sha256':h})
    docs={k:document(pins,k,v) for k,v in c['documents'].items()}
    require(output.parent.is_relative_to(pins.path(c['edition'])/'work') or output.parent.is_relative_to(Path(__file__).resolve().parent),'QA output outside edition work or owned helper directory')
    targets=[]
    for d in c['documents'].values():
        if 'deliverable_docx' in d:
            target=pins.path(d['deliverable_docx'])
            require(target.parent==pins.path(c['edition'])/'deliverables','DOCX alias scope')
            pins.pin({'path':str(target),'sha256':d['docx']['sha256']})
    if copy_pdfs:
        for kind,d in c['documents'].items():
            source=pins.pin(d['pdf']);target=pins.path(d['deliverable_pdf'])
            require(target.parent==pins.path(c['edition'])/'deliverables','Delivery scope')
            require(not target.exists() or sha(target)==sha(source),'Refuse differing delivered PDF')
            targets.append((source,target))
    pins.reread();require(Path(config_path).read_bytes()==raw,'Config mutation')
    # Validation is now complete; publication uses exclusive creation only.
    for source,target in targets:
        if not target.exists():
            with target.open('xb') as f:f.write(source.read_bytes())
        pins.pin({'path':str(target),'sha256':sha(source)})
    pins.reread()
    edition=pins.path(c['edition'])
    relative=lambda p:Path(os.path.relpath(p,edition)).as_posix()
    reviews=sorted({relative(pins.path(d['visual'][page][key]['path'])) for d in c['documents'].values() for page in d['visual'] for key in ['report','ledger']})
    result={'status':'PASS_TECHNICAL_PREPARATION','scope':'Internal computational QA only; no author approval, rights clearance, submission, or physical accuracy certification',
      'created_utc':datetime.now(timezone.utc).isoformat(),'config_sha256':hashlib.sha256(raw).hexdigest(),
      'helper_sha256':sha(Path(__file__)),'documents':docs,'old_scientific_files':563,'old_total_files':len(data['files']),
      'files':{relative(p):h for p,h in pins.files.items()},'review_files':reviews,
      'config_path':relative(Path(config_path).resolve()),'copied_or_identical_pdfs':[relative(t) for _,t in targets]}
    with output.open('x') as f:json.dump(result,f,indent=2);f.write('\n')
    return result
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('config',type=Path);ap.add_argument('--output',required=True,type=Path);ap.add_argument('--copy-pdfs',action='store_true');a=ap.parse_args()
    print(json.dumps(run(a.config,a.output,a.copy_pdfs),indent=2))
