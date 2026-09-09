"""Independent saved final-QA authentication, no finalizer/scientific execution."""
from collections import Counter
from datetime import datetime,timezone
import hashlib,json,os,re,time,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

HERE=Path(__file__).resolve().parent
WORK=HERE.parents[2]
PROJECT=WORK.parents[2]
EDITION=PROJECT/'submission_revision_20260908/renewed_manuscript_v9'
TOOLS=WORK/'independent_environment/final_qa_tools'

def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(2**20),b''):h.update(b)
    return h.hexdigest()
def pointer(d,s):
    for k in s.strip('/').split('/') if s else []:
        k=k.replace('~1','/').replace('~0','~');d=d[int(k)] if isinstance(d,list) else d[k]
    return d
def main():
    started=time.monotonic();pins={};json_reads={}
    def locate(p,base=PROJECT):
        p=Path(p);return p if p.is_absolute() else base/p
    def pin(p,h):
        p=locate(p)
        assert not any(x.is_symlink() for x in [p,*p.parents]),str(p)
        assert p.is_file() and sha(p)==h,str(p)
        key=str(p.resolve());assert key not in pins or pins[key]==h
        pins[key]=h;return p
    def read(p,h):
        p=pin(p,h);raw=p.read_bytes();assert hashlib.sha256(raw).hexdigest()==h
        json_reads[str(p.resolve())]=h;return json.loads(raw)
    qpath=EDITION/'work/FINAL_TECHNICAL_QA_v9.json'
    q=read(qpath,'24282221edca95bfefe05e2add3f5974e49c406d6d68565d421f5faabb1bcc26')
    assert q['status']=='PASS_TECHNICAL_PREPARATION' and len(q['files'])==3635
    c=read(EDITION/q['config_path'],'419e75ceec4cf80b364572f0e776a40f2d6506a90178f79b794ad8b42a510642')
    assert c['finalization_authorized'] is True and q['config_sha256']==sha(EDITION/q['config_path'])
    pin(TOOLS/'finalize_v9_v2.py',q['helper_sha256'])
    assert q['helper_sha256']=='027a192786e727344b45e7de5f7eb4e158abe85222d856e40afdab6ba8896341'
    ext=c['external_readonly_pins'];assert ext==q['external_readonly_files'] and len(ext)==360
    proposal=read(TOOLS/'V9_EXTERNAL_READONLY_PROPOSAL.json','148540058cbf4ffefaed38f197e1a939d5ffb720a29f360cb92b921571eb3728')
    assert ext==proposal['external_readonly_pins']
    root_inventory={}
    for rel,h in q['files'].items():
        p=pin(EDITION/rel,h);canonical=str(p.resolve())
        assert canonical not in root_inventory,'Alias inventory duplicate'
        root_inventory[canonical]=h
        if not p.resolve().is_relative_to(PROJECT):assert ext.get(canonical)==h
    actual_external={p:h for p,h in root_inventory.items() if not Path(p).is_relative_to(PROJECT)}
    assert actual_external==ext
    assert Counter(Path(p).suffix for p in ext)=={'.py':350,'.npz':10}
    document_results={};visual_total=[]
    for kind,dc in c['documents'].items():
        saved=q['documents'][kind]
        paths={k:pin(v['path'],v['sha256']) for k,v in saved['artifacts'].items()}
        for key in ('source','builder','build','layout','page_map','docx','pdf'):
            assert paths[key].resolve()==locate(dc[key]['path']).resolve() and sha(paths[key])==dc[key]['sha256']
        build=read(paths['build'],dc['build']['sha256']);layout=read(paths['layout'],dc['layout']['sha256']);page_map=read(paths['page_map'],dc['page_map']['sha256'])
        assert build['source_sha256']==dc['source']['sha256'] and build['builder_sha256']==dc['builder']['sha256'] and build['docx_sha256']==dc['docx']['sha256']
        assert layout['status']=='PASS_PROGRAMMATIC' and not layout['violations'] and layout['pdf_sha256']==dc['pdf']['sha256']
        doc=Document(paths['docx']);pdf=PdfReader(paths['pdf'])
        assert len(pdf.pages)==dc['pages']==saved['pages']==layout['pages']
        assert not doc.core_properties.author and not doc.core_properties.last_modified_by
        assert not (pdf.metadata or {}).get('/Author','')
        numbered=[p for p in doc.paragraphs if p._p.xpath('./m:oMath') and re.search(r'\((?:S)?\d+\)$',p.text)]
        actual=[len(numbered),len(doc.element.xpath('.//m:oMath')),len(doc.tables),len(doc.inline_shapes)]
        assert actual==saved['counts']==c['content_contract'][kind] and build['equations']==actual[0]
        for table in doc.tables:
            props=table.rows[0]._tr.find(qn('w:trPr'));assert props is not None
            header=props.find(qn('w:tblHeader'));assert header is not None and header.get(qn('w:val'),'true') not in {'0','false','off'}
        for tag in ['w:ins','w:del','w:moveFrom','w:moveTo','w:commentReference','w:commentRangeStart','w:commentRangeEnd']:
            assert not list(doc.element.iter(qn(tag)))
        with zipfile.ZipFile(paths['docx']) as z:
            comments=[n for n in z.namelist() if 'comment' in n.lower()]
            assert len(comments)<=1 and all(n=='word/comments.xml' for n in comments)
            for n in comments:
                element=ET.fromstring(z.read(n));assert element.tag==qn('w:comments') and len(element)==0 and not (element.text or '').strip()
        reference_numbers=[int(m[1]) for p in doc.paragraphs if (m:=re.match(r'^\[(\d+)\]',p.text))]
        assert reference_numbers==list(range(1,24))
        oldspec=c['v8_sources'][kind];oldpath=pin(oldspec['path'],oldspec['sha256'])
        entries=lambda txt:re.findall(r'^\[(\d+)\] ([^\n]+)$',txt,re.M)
        assert entries(paths['source'].read_text())[:21]==entries(oldpath.read_text())
        toc=[]
        if kind=='supplement':
            source=paths['source'].read_bytes();token=b'NAVIGATION_PLACEHOLDER'
            assert source.count(token)==1 and b'\n'+token+b'\n' in source
            headings=re.findall(r'^# (S[0-9]+ [^\n]+|References)[ \t]*$',source.decode(),re.M)
            assert len(headings)==14 and set(headings)==set(dc['toc_labels'])
            rows=['| Section | PDF page |','| --- | ---: |']
            for heading in headings:
                label=dc['toc_labels'][heading];page=page_map[heading]
                assert any(p.text==label+'\t'+str(page) for p in doc.paragraphs)
                text=pdf.pages[page-1].extract_text() or ''
                assert re.sub(r'\s+',' ',heading) in re.sub(r'\s+',' ',text)
                rows.append('| '+label+' | '+str(page)+' |');toc.append(dict(heading=heading,label=label,page=page))
            assert paths['reading_source'].read_bytes()==source.replace(token,'\n'.join(rows).encode())
        assert set(dc['visual'])=={str(i) for i in range(1,dc['pages']+1)}
        assert len(saved['visual'])==dc['pages']
        bypage={r['page']:r for r in saved['visual']}
        observations=[]
        for row in layout['ledger']:
            page=row['page'];spec=dc['visual'][str(page)];v=bypage[page]
            current=pin(paths['pdf'].parent/f'page-{page}.png',row['png_sha256'])
            witness=read(spec['ledger']['path'],spec['ledger']['sha256']);r=pointer(witness,spec['record_pointer'])
            reviewed=pin(spec['reviewed_png']['path'],spec['reviewed_png']['sha256'])
            reviewedpdf=pin(spec['reviewed_pdf']['path'],spec['reviewed_pdf']['sha256'])
            pin(spec['report']['path'],spec['report']['sha256'])
            assert r['page']==spec['reviewed_page']==v['reviewed_page']
            assert r['result']=='PASS' and r['individually_viewed_original_resolution'] is True
            assert r['png_sha256']==sha(reviewed)==sha(current)==v['png_sha256']
            assert witness['pdf_sha256']==sha(reviewedpdf)==v['reviewed_pdf_sha256']
            assert spec['ledger']['sha256']==v['ledger_sha256'] and spec['report']['sha256']==v['review_report_sha256']
            observation=dict(kind=kind,page=page,reviewed_page=r['page'],png_sha256=v['png_sha256'],
                witness_same_document=(sha(reviewedpdf)==dc['pdf']['sha256']),fresh_visual_inspection=False)
            observations.append(observation);visual_total.append(observation)
        aliases={}
        for alias,sourcekey in [('deliverable_docx','docx'),('deliverable_pdf','pdf')]:
            if alias in dc:
                ap=pin(dc[alias],dc[sourcekey]['sha256']);aliases[alias]=dict(path=str(ap),sha256=sha(ap))
        document_results[kind]=dict(pages=len(pdf.pages),native_counts=actual,references=len(reference_numbers),
                                   toc=toc,visual=observations,aliases=aliases,author_metadata_blank=True)
    assert len(visual_total)==134
    assert q['old_scientific_files']==563 and q['old_total_files']==754
    result=dict(status='PASS_SAVED_ACTUAL_FINAL_QA_AUDIT',utc=datetime.now(timezone.utc).isoformat(),
        source_sha256=sha(Path(__file__)),qa_sha256=sha(qpath),config_sha256=q['config_sha256'],helper_sha256=q['helper_sha256'],
        finalizer_created_utc=q['created_utc'],finalizer_inventory_count=len(root_inventory),internal_inventory_count=3275,
        external_inventory_count=len(ext),external_extensions=dict(Counter(Path(p).suffix for p in ext)),
        inventory=root_inventory,auxiliary_pins={p:h for p,h in pins.items() if p not in root_inventory},
        metadata_json_reads=json_reads,documents=document_results,
        visual_existing_witness_count=134,visual_same_document=sum(v['witness_same_document'] for v in visual_total),
        visual_full_png_inherited=sum(not v['witness_same_document'] for v in visual_total),
        old_scientific_files=563,old_total_files=754,elapsed_audit_seconds=time.monotonic()-started,
        scientific_payloads_parsed=0,finalizers_rerun=0,new_visual_inspections=0,
        limitations=['Saved visual witnesses reauthenticated, not a fresh visual inspection.',
                     'All scientific NPZ/archive pins were hash-only; no member materialization.',
                     'Technical preparation is not human author approval, rights clearance, or journal submission.'])
    with (HERE/'QA.json').open('x') as f:json.dump(result,f,indent=2,sort_keys=True);f.write('\n')
    print(json.dumps(dict(status=result['status'],qa_sha256=sha(HERE/'QA.json'),inventory=len(root_inventory),visual134=len(visual_total),
        same_document=result['visual_same_document'],inherited=result['visual_full_png_inherited'],seconds=result['elapsed_audit_seconds'])))

if __name__=='__main__':main()
