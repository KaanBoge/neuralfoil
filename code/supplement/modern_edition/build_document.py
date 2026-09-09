"""Style the portable Pandoc DOCX as a modern research reading edition.

Run with the document-runtime Python. Scientific content is assembled separately;
this module changes typography, table geometry, navigation, and OOXML semantics.
"""
from __future__ import annotations

import argparse
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
BODY = "Noto Serif"
SANS = "Carlito"
MATH = "DejaVu Math TeX Gyre"
WIDTH = 6.7
TITLE = "NeuralFoil Accuracy and Reliability Against Wind Tunnel Measurements"


def node(tag, **attrs):
    e = OxmlElement(tag)
    for key, value in attrs.items():
        e.set(qn(key), str(value))
    return e


def font(run, name=BODY, size=None, color="17212B", bold=None):
    run.font.name = name
    rp = run._r.get_or_add_rPr()
    rf = rp.find(qn("w:rFonts"))
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{attr}"), name)
    for attr in list(rf.attrib):
        if attr.endswith("Theme"):
            del rf.attrib[attr]
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def style(doc, name, family, size, before=0, after=0, bold=False):
    st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = family
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    rf = st.element.get_or_add_rPr().find(qn("w:rFonts"))
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{attr}"), family)
    for attr in list(rf.attrib):
        if attr.endswith("Theme"):
            del rf.attrib[attr]
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.widow_control = True
    pf.line_spacing = 1.2
    return st


def bookmark(p, name, ident):
    p._p.insert(0, node("w:bookmarkStart", **{"w:id": ident, "w:name": name}))
    p._p.append(node("w:bookmarkEnd", **{"w:id": ident}))


def hyperlink(p, label, anchor, size=10.5):
    h = node("w:hyperlink", **{"w:anchor": anchor, "w:history": "1"})
    r = node("w:r")
    rp = node("w:rPr")
    rp.append(node("w:rFonts", **{"w:ascii": SANS, "w:hAnsi": SANS}))
    rp.append(node("w:sz", **{"w:val": int(size * 2)}))
    rp.append(node("w:color", **{"w:val": "24586A"}))
    r.append(rp)
    t = node("w:t"); t.text = label; r.append(t)
    h.append(r); p._p.append(h)


def page_field(p):
    r = p.add_run()
    r._r.append(node("w:fldChar", **{"w:fldCharType": "begin"}))
    t = node("w:instrText", **{"xml:space": "preserve"}); t.text = " PAGE "
    r._r.append(t)
    r._r.append(node("w:fldChar", **{"w:fldCharType": "end"}))
    font(r, SANS, 9, "000000")


def equation(table):
    label = table.cell(0, 1).text.strip()
    xml = node("w:p"); table._tbl.addprevious(xml)
    p = Paragraph(xml, table._parent)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(WIDTH / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    pf.line_spacing = 1.15
    pf.space_before = Pt(8); pf.space_after = Pt(9)
    pf.keep_together = True
    p.add_run("\t")
    for child in table.cell(0, 0).paragraphs[0]._p:
        if child.tag != qn("w:pPr"):
            p._p.append(deepcopy(child))
    font(p.add_run(f"\t{label}"), SANS, 10)
    table._tbl.getparent().remove(table._tbl)


def table_widths(table, caption):
    mapping = {
        "1": [1.5, 1.65, 1.4, 2.15],
        "2": [2.5, .65, 3.55],
        "3": [2.65, 1.1, 1.1, 1.85],
        "4": [1.05, 1.4, 1.25, 1.45, 1.55],
        "5": [1.7, 1.6, 1.8, 1.6],
        "6": [2.2, 1.9, 1.35, 1.25],
        "7": [2.65, .85, .95, 2.25],
        "8": [1.5, .55, .72, 2.12, 1.81],
        "A1": [2.4, 2.05, 2.25],
        "C1": [1.1, .8, .75, 2.05, 2.0],
        "C2": [1.55, .9, .9, .9, 1.0, 1.45],
        "C3": [1.2, 1.1, 1.4, 1.45, 1.55],
        "C4": [1.65, 1.05, 1.15, 1.15, 1.7],
        "C5": [2.05, 2.2, 2.45],
        "C6": [1.7, 1.1, 2.0, 1.9],
        "D1": [1.5, 1.05, 1.95, 2.2],
        "E1": [1.65, 1.0, 4.05],
        "F1": [1.5, 5.2],
    }
    widths = mapping.get(caption)
    if widths is None:
        widths = [1.15, 5.55] if len(table.columns) == 2 else [WIDTH / len(table.columns)] * len(table.columns)
    assert len(widths) == len(table.columns), (caption, len(widths), len(table.columns))
    assert abs(sum(widths) - WIDTH) < .001
    return widths


def format_table(table, caption):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = table_widths(table, caption)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    pr = table._tbl.tblPr
    borders = pr.find(qn("w:tblBorders"))
    if borders is not None: pr.remove(borders)
    borders = node("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        borders.append(node(f"w:{side}", **{"w:val":"single", "w:sz":"4", "w:color":"D9D9D9"}))
    pr.append(borders)
    for idx, row in enumerate(table.rows):
        tr = row._tr.get_or_add_trPr()
        tr.append(node("w:cantSplit"))
        if idx == 0: tr.append(node("w:tblHeader", **{"w:val":"true"}))
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            shd = tc.find(qn("w:shd"))
            if shd is not None: tc.remove(shd)
            tc.append(node("w:shd", **{"w:fill": "EAF0F3" if idx == 0 else ("F8FAFB" if idx % 2 == 0 else "FFFFFF")}))
            mar = node("w:tcMar")
            for side, val in (("top", 80), ("bottom", 80), ("start", 85), ("end", 85)):
                mar.append(node(f"w:{side}", **{"w:w":val, "w:type":"dxa"}))
            tc.append(mar)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.12
                pf.keep_together = True
                # Compact comparison tables should be read as a unit; longer
                # evidence matrices may continue with their repeated header.
                compact=caption in {"1","2","3","4","5","6","7","8","C1","C2","C3","C4","C5","C6","F1"}
                pf.keep_with_next = (idx < len(table.rows)-1) if compact else idx == 0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if idx > 0 and col_idx > 0 and re.fullmatch(r"[+−\-\d, .%°^a]+", p.text.strip()):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    is_hash = bool(re.search(r"[0-9a-f]{50,}", run.text))
                    font(run, "Liberation Mono" if is_hash else SANS,
                         8.2 if is_hash else (9.5 if len(widths) >= 5 else 10),
                         "000000" if idx == 0 else "17212B", bold=True if idx == 0 else None)
    following=table._tbl.getnext()
    if following is not None and following.tag==qn("w:p"):
        after=Paragraph(following,table._parent)
        existing=after.paragraph_format.space_before
        after.paragraph_format.space_before=Pt(max(7,existing.pt if existing else 0))
        if after.text.startswith("a Sequential estimate"):
            after.paragraph_format.line_spacing=1.12
            for run in after.runs: font(run,SANS,9.5)


def format_document(doc, page_map):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(.9)
    sec.top_margin = Inches(.78); sec.bottom_margin = Inches(.75)
    sec.header_distance = Inches(.32); sec.footer_distance = Inches(.32)
    sec.different_first_page_header_footer = True
    base = style(doc, "Normal", BODY, 11, after=6)
    base.paragraph_format.line_spacing = 1.22
    for name in ("Body Text", "First Paragraph"):
        st = style(doc, name, BODY, 11, after=6)
        st.base_style = base
        st.paragraph_format.line_spacing = 1.22
        st.paragraph_format.first_line_indent = Inches(0)
    st = style(doc, "Title", SANS, 28, after=15, bold=True)
    st.paragraph_format.line_spacing = 1.02
    st.paragraph_format.keep_with_next = True
    style(doc, "Subtitle", SANS, 12.5, after=10)
    for name, size, before, after in (("Heading 1",17,18,9), ("Heading 2",13,15,7), ("Heading 3",11.5,11,5)):
        st = style(doc,name,SANS,size,before,after,True)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True
        st.paragraph_format.line_spacing = 1.12
    style(doc,"Caption",SANS,10,4,10)
    style(doc,"Compact",SANS,10,0,4)
    style(doc,"Source Code","Liberation Mono",9.2,0,5)
    for st in doc.styles:
        if st.type == WD_STYLE_TYPE.CHARACTER and st.name in {"Verbatim Char", "Source Code"}:
            st.font.name = "Liberation Mono"; st.font.size = Pt(9.2)
    # Explicit default language and a math font that is present in the renderer.
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        for lang in defaults.iter(qn("w:lang")):
            lang.set(qn("w:val"), "en-US")
    mathpr = doc.settings.element.find(qn("m:mathPr"))
    if mathpr is None: mathpr = node("m:mathPr"); doc.settings.element.append(mathpr)
    old = mathpr.find(qn("m:mathFont"))
    if old is not None: mathpr.remove(old)
    mathpr.insert(0,node("m:mathFont", **{"m:val":MATH}))
    # Pandoc emits U+203E as a generic accent; LibreOffice renders it as an
    # acute-like glyph. A native top bar preserves the intended mean notation.
    for accent in list(doc.element.iter(qn("m:acc"))):
        char=accent.find(f"{qn('m:accPr')}/{qn('m:chr')}")
        if char is not None and char.get(qn("m:val")) in {"‾", "̄"}:
            bar=node("m:bar")
            prop=node("m:barPr"); prop.append(node("m:pos", **{"m:val":"top"}))
            bar.append(prop)
            expression=accent.find(qn("m:e"))
            bar.append(deepcopy(expression))
            accent.getparent().replace(accent,bar)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(hp.add_run("NEURALFOIL ACCURACY AND RELIABILITY"),SANS,8.2,"000000")
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
    font(fp.add_run("Kaan Boge"),SANS,8.5,"000000")
    fp.add_run("\t"); page_field(fp)
    sec.first_page_header.paragraphs[0].text = ""
    first_footer = sec.first_page_footer.paragraphs[0]
    first_footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; page_field(first_footer)

    heading_info=[]; counter=8000
    table_caption=""; before_abstract=True; reference_mode=False
    for p in list(doc.paragraphs):
        text=p.text.strip(); pf=p.paragraph_format
        pf.widow_control=True
        if p.style.name=="Title":
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs: font(r,SANS,28,"000000",True)
        elif before_abstract and text:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            pf.space_after=Pt(9); pf.line_spacing=1.1
            for r in p.runs: font(r,SANS,11)
        if text=="Abstract": before_abstract=False
        if p.style.name.startswith("Heading"):
            for r in p.runs: font(r,SANS,None,"000000",True)
            for b in list(p._p.xpath("./w:pPr/w:pBdr")): b.getparent().remove(b)
            anchor=f"section_{counter}"; bookmark(p,anchor,counter); counter+=1
            heading_info.append({"text":text,"anchor":anchor,"level":int(p.style.name[-1])})
            if text in {"Reader guide", "1 Introduction", "Appendix H Notation and concepts", "References"}:
                pf.page_break_before=True
            if text in {"Abstract", "Reader guide", "Contents"}:
                pf.space_before=Pt(11)
        if text=="References": reference_mode=True
        elif reference_mode and text.startswith("["):
            pf.left_indent=Inches(.25); pf.first_line_indent=Inches(-.25)
            pf.space_after=Pt(9); pf.line_spacing=1.12
            for r in p.runs: font(r,BODY,10)
        if text.startswith("Keywords:"):
            pf.line_spacing=1.05; pf.space_before=Pt(7); pf.space_after=Pt(11)
            for r in p.runs: font(r,SANS,9.5)
        if re.match(r"^Table\s+(?:[A-Z]\d+|\d+)\.",text):
            pf.keep_with_next=True; pf.keep_together=True
            pf.space_before=Pt(12); pf.space_after=Pt(5); pf.line_spacing=1.12
            for r in p.runs: font(r,SANS,10,"000000",True)
        if p._p.xpath(".//w:drawing"):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pf.keep_together=True; pf.keep_with_next=True
            pf.space_before=Pt(7); pf.space_after=Pt(3); pf.line_spacing=1.0
        elif p.style.name=="Caption" or re.match(r"^Figure \d+\.",text):
            p.style=doc.styles["Caption"]
            pf.keep_together=True; pf.keep_with_next=False
            pf.line_spacing=1.12; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs: font(r,SANS,10)

    # Reformat native display equations and genuine data tables independently.
    caption_by_table={}
    current=""
    for child in doc.element.body:
        if child.tag==qn("w:p"):
            txt="".join(child.xpath(".//w:t/text()"))
            m=re.match(r"^Table\s+([A-Z]?\d+)\.",txt)
            if m: current=m.group(1)
            elif child.xpath("./w:pPr/w:pStyle") and "Appendix H" in txt: current=""
        elif child.tag==qn("w:tbl"):
            caption_by_table[child]=current
    for table in list(doc.tables):
        if len(table.rows)==1 and len(table.columns)==2 and re.fullmatch(r"\(\d+[a-z]?\)",table.cell(0,1).text.strip()):
            equation(table)
        else:
            cap=caption_by_table.get(table._tbl,"")
            if table.cell(0,0).text.strip()=="Symbol": cap=""
            format_table(table,cap)

    # Replace the reserved navigation marker with meaningful internal links.
    marker=next(p for p in doc.paragraphs if p.text.strip()=="NAVIGATION_PLACEHOLDER")
    selected=[]
    for h in heading_info:
        if h["text"].startswith(("1 ","2 ","3 ","4 ","5 ","6 ","Appendix ")) or h["text"]=="References":
            if h["level"]==1: selected.append(h)
    for h in selected:
        p=marker.insert_paragraph_before()
        p.paragraph_format.line_spacing=1.0
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        hyperlink(p,h["text"],h["anchor"],10.5)
        value=str(page_map.get(h["text"],"  "))
        font(p.add_run("\t"+value),SANS,10.5)
    marker._p.getparent().remove(marker._p)
    props=doc.core_properties
    props.title=TITLE; props.author="Kaan Boge"
    props.subject="Measurement audit of airfoil predictions with full methods and reproducibility appendices"
    props.keywords="NeuralFoil, XFOIL, airfoil, wind tunnel, model validation, uncertainty"
    props.comments=""; props.last_modified_by="Kaan Boge"
    return heading_info


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("input",type=Path)
    parser.add_argument("output",type=Path)
    parser.add_argument("--page-map",type=Path)
    args=parser.parse_args()
    pages=json.loads(args.page_map.read_text()) if args.page_map else {}
    doc=Document(args.input)
    headings=format_document(doc,pages)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(args.output)
    (ROOT/"work"/"headings.json").write_text(json.dumps(headings,indent=2)+"\n")
    print(args.output)


if __name__=="__main__": main()
